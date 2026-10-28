"""Fills the official SIH 2026 idea template for PS SIH26263 with a visual-first deck.

Keeps the template frame (logo, team-name oval, headings, footer, pointers) and adds
diagrams, flowcharts and prototype screenshots in the content area.

Usage (from ppt/):
  node export_icons.js .build/icons
  python build_template_deck.py --team-id <ID> --live-url <https://...> --out SIH26263_AI_Riders_Idea.pptx
"""
import argparse
import json
from pathlib import Path

import qrcode
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
SHOTS = ROOT / "docs" / "screenshots"
ICONS = HERE / ".build" / "icons"
BENCH = json.loads((ROOT / "docs" / "benchmark.json").read_text())
API_TESTS = sum(f.read_text().count("\ndef test_") for f in (ROOT / "backend" / "tests").glob("test_*.py"))
HELD = BENCH["identity_matching_labelled_pairs"]["held_out"]
# Real interview evidence, filled by the team after talking to an admission cell (docs/validation/INTERVIEW_KIT.md).
# Shown only when a quote exists and the speaker agreed to be quoted; nothing is shown otherwise.
_FV = HERE / "field_validation.json"
FIELD = json.loads(_FV.read_text()) if _FV.exists() else {}
HAS_QUOTE = bool(FIELD.get("quote", "").strip()) and FIELD.get("permission_to_quote") is True
COVERAGE = json.loads((ROOT / "docs" / "evidence.json").read_text())["coverage_percent"]
RULE_COUNT = len(__import__("re").findall(r'"[A-Z]{3}-\d{3}": \("', (ROOT / "backend" / "app" / "validation.py").read_text()))

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x00, 0x70, 0xC0)
ORANGE = RGBColor(0xE4, 0x6C, 0x0A)
GREEN = RGBColor(0x1E, 0x7B, 0x3E)
RED = RGBColor(0xB4, 0x2B, 0x2B)
INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x59, 0x59, 0x59)
LINE = RGBColor(0xC9, 0xD3, 0xE3)
PALE = RGBColor(0xEE, 0xF3, 0xFA)
PALE_RED = RGBColor(0xFC, 0xEC, 0xEA)
PALE_GREEN = RGBColor(0xE8, 0xF5, 0xEC)
PALE_ORANGE = RGBColor(0xFD, 0xF1, 0xE6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

PS_TITLE = ("Development of an APAAR ID-Based Student Data Management and Verification Platform "
            "for AICTE-Approved Institutions")


# --------------------------------------------------------------------------- primitives
def box(slide, x, y, w, h, fill=None, line=None, radius=0.12, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=0.75):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def write(tf, parts, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, spacing=None):
    """parts: str, or list of paragraphs; each paragraph a str or list of (text, {opts}) runs."""
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    paras = [parts] if isinstance(parts, str) else parts
    tf.text = ""
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        runs = [(para, {})] if isinstance(para, str) else para
        for text, opts in runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = FONT
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", False)
            f.color.rgb = opts.get("color", color)
            if opts.get("link"):
                r.hyperlink.address = opts["link"]


def text(slide, x, y, w, h, parts, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    write(tb.text_frame, parts, **kw)
    return tb


def label_box(slide, x, y, w, h, parts, fill, line=None, color=INK, size=12, bold=False, align=PP_ALIGN.CENTER, radius=0.1):
    s = box(slide, x, y, w, h, fill=fill, line=line, radius=radius)
    write(s.text_frame, parts, size=size, color=color, bold=bold, align=align)
    return s


def pointer(slide, x, y, w, txt, h=0.34):
    """Template pointer text, kept verbatim, styled as a section tab."""
    return label_box(slide, x, y, w, h, [[("❖ ", {"color": ORANGE}), (txt, {})]], fill=NAVY, color=WHITE,
                     size=12, bold=True, align=PP_ALIGN.LEFT, radius=0.08)


def icon(slide, name, x, y, d, bg=NAVY, variant="white"):
    if bg is not None:
        box(slide, x, y, d, d, fill=bg, shape=MSO_SHAPE.OVAL)
    pad = d * 0.22 if bg is not None else 0
    slide.shapes.add_picture(str(ICONS / f"{name}_{variant}.png"), Inches(x + pad), Inches(y + pad),
                             Inches(d - 2 * pad), Inches(d - 2 * pad))


def arrow(slide, x1, y1, x2, y2, color=GREY, width=1.5, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if head:
        ln = c.line._get_or_add_ln()
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    return c


def picture(slide, path, x, y, w, crop=(0, 0, 0, 0), border=True):
    """crop = (left, top, right, bottom) fractions. Height follows the cropped aspect ratio."""
    with Image.open(path) as im:
        iw, ih = im.size
    l, t, r, b = crop
    aspect = (iw * (1 - l - r)) / (ih * (1 - t - b))
    h = w / aspect
    if border:
        box(slide, x - 0.03, y - 0.03, w + 0.06, h + 0.06, fill=WHITE, line=LINE, radius=0.04)
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = l, t, r, b
    return h


def framed(slide, path, x, y, w, h, region=(0, 0, 0, 0)):
    """Fill a w x h frame from the region (left, top, right, bottom crop fractions), trimming its bottom or
    sides so the screenshot keeps its proportions."""
    with Image.open(path) as im:
        iw, ih = im.size
    l, t, r, b = region
    rw, rh = iw * (1 - l - r), ih * (1 - t - b)
    if rw / rh < w / h:
        b += (rh - rw * h / w) / ih
    else:
        trim = (rw - rh * w / h) / iw / 2
        l, r = l + trim, r + trim
    box(slide, x - 0.03, y - 0.03, w + 0.06, h + 0.06, fill=WHITE, line=LINE, radius=0.04)
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = l, t, r, b


def qr(slide, data, x, y, d, path):
    img = qrcode.make(data, box_size=10, border=1)
    img.save(path)
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(d), Inches(d))


def remove(shape):
    shape._element.getparent().remove(shape._element)


def set_runs(shape, txt, size=None, color=None):
    """Replace a shape's text, keeping the first run's formatting."""
    tf = shape.text_frame
    first_p = next(p for p in tf.paragraphs if p.runs)
    keep = first_p.runs[0]
    for p in list(tf.paragraphs):
        if p._p is not first_p._p:
            p._p.getparent().remove(p._p)
    for r in first_p.runs[1:]:
        r._r.getparent().remove(r._r)
    for br in first_p._p.findall(qn("a:br")):  # template titles carry trailing line breaks
        first_p._p.remove(br)
    keep.text = txt
    if size:
        keep.font.size = Pt(size)
    if color:
        keep.font.color.rgb = color


def find(slide, prefix):
    return next(s for s in slide.shapes if s.name.startswith(prefix))


# --------------------------------------------------------------------------- slides
def bare(url: str) -> str:
    return url.split("://", 1)[-1].rstrip("/")


def slide1(s, team_id, team_name, build, live_url="", video_url=""):
    tb = find(s, "TextBox 9")
    values = {"Problem Statement ID": "SIH26263", "Problem Statement Title": PS_TITLE, "Theme": "Smart Automation",
              "PS Category": "Software", "Team ID": team_id, "Team Name": team_name}
    for p in tb.text_frame.paragraphs:
        if not p.runs:
            continue
        label = p.runs[0].text
        key = next((k for k in values if label.startswith(k)), None)
        if key is None:
            continue
        if key == "PS Category":
            p.runs[0].text = "PS Category- "
        for r in p.runs:
            r.font.size = Pt(17)
        v = p.add_run()
        v.text = (" " if not p.runs[0].text.endswith(" ") else "") + values[key]
        v.font.size = Pt(15 if key == "Problem Statement Title" else 17)
        v.font.bold = False
        v.font.name = FONT
        v.font.color.rgb = BLUE
        p.space_after = Pt(10)
        p.space_before = Pt(0)
        p.line_spacing = 1.0
        p.alignment = PP_ALIGN.LEFT
    tb.width = Inches(6.2)
    tb.top = Inches(2.45)
    # The template title runs under the SIH logo when Garamond is substituted; keep it clear of the logo.
    title = find(s, "Title 7")
    title.width = Inches(10.2)
    for p in title.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(36)
    # The template's large SIH illustration and its hexagon backdrop give way to the working prototype itself;
    # the SIH logo in the corner (template frame) stays.
    remove(find(s, "Picture 4"))
    remove(find(s, "Freeform"))
    heading = find(s, "Subtitle 3")  # "TITLE PAGE", centred over the details column so the right side is free
    heading.left, heading.width = Inches(0.36), Inches(6.4)

    box(s, 0.45, 5.92, 6.3, 1.3, fill=PALE, line=LINE)
    icon(s, "LuShieldCheck", 0.58, 6.0, 0.62, bg=NAVY)
    lines = [[("VIDYA PRAMAAN", {"bold": True, "size": 18, "color": NAVY})],
             [("One student. One APAAR ID. One verifiable record.", {"size": 11, "color": INK})],
             [("Live: ", {"size": 10.5, "color": GREY}),
              (bare(live_url) if live_url else "./run.sh", {"size": 10.5, "color": BLUE, "link": live_url or None})]]
    if video_url:
        lines.append([("Demo video (5 min): ", {"size": 10.5, "color": GREY}),
                      (bare(video_url), {"size": 10.5, "color": BLUE, "link": video_url})])
    lines.append([(f"7 roles · {RULE_COUNT} rules · {API_TESTS} tests · synthetic data", {"size": 10.5, "color": GREY})])
    text(s, 1.3, 5.95, 4.15, 1.24, lines, size=11)
    if live_url:
        qr(s, live_url, 5.55, 6.02, 1.1, build / "qr_live1.png")

    # The prototype, as a judge would see it
    text(s, 7.0, 1.6, 6.0, 0.3, [[("Working prototype", {"bold": True, "color": NAVY}),
                                  ("  ·  live now, every screen below is real", {"color": GREY, "size": 11})]], size=13)
    framed(s, SHOTS / "10_aicte_dashboard.png", 7.0, 1.95, 4.5, 3.02, region=(0.178, 0.058, 0, 0))
    framed(s, SHOTS / "30_mobile_dashboard.png", 11.62, 1.95, 1.38, 3.02)
    cap = dict(size=10, color=GREY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
    text(s, 7.0, 5.0, 4.5, 0.26, "AICTE national view: every breach named with its rule", **cap)
    text(s, 11.5, 5.0, 1.6, 0.26, "On a phone (PWA)", **cap)
    framed(s, SHOTS / "14_review_identity_duplicate.png", 7.0, 5.36, 4.0, 1.45, region=(0.54, 0.015, 0, 0))
    framed(s, SHOTS / "19_review_shared_contact_cluster.png", 11.12, 5.36, 1.88, 1.45, region=(0.6, 0.165, 0.07, 0.425))
    text(s, 7.0, 6.86, 4.0, 0.26, "Same person under two APAAR IDs: reasons shown", **cap)
    text(s, 10.9, 6.86, 2.3, 0.26, "1 phone, 8 students, 2 colleges", **cap)


def slide2(s, team_name):
    title = find(s, "Title 1")
    set_runs(title, "VIDYA PRAMAAN", size=30)
    title.top, title.height = Inches(0.3), Inches(0.55)
    tf = title.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = tf.margin_bottom = Inches(0)
    text(s, 2.2, 0.8, 8.9, 0.34, [[("29.6 lakh", {"bold": True, "color": NAVY}),
                                    (" enrolments in AICTE institutions (2025-26)  ·  ", {"color": GREY}),
                                    ("24.25 crore", {"bold": True, "color": NAVY}),
                                    (" APAAR IDs issued (Mar 2026)", {"color": GREY})]],
         size=12, align=PP_ALIGN.CENTER)
    remove(find(s, "TextBox 8"))
    text(s, 0.3, 1.2, 12.7, 0.38, [[("❖ ", {"color": BLUE}), ("Proposed Solution (Describe your Idea/Solution/Prototype)", {})]],
         size=17, bold=True, color=NAVY)

    # Detailed explanation: the end-to-end flow the PS asks for
    pointer(s, 0.35, 1.66, 4.6, "Detailed explanation of the proposed solution")
    steps = [("LuUpload", "Institution upload", "Excel · CSV · ERP API"), ("LuFingerprint", "APAAR validation", "registry + consent"),
             ("LuListChecks", "Data verification", f"{RULE_COUNT} explained rules"), ("LuDatabase", "Central repository", "one record per student"),
             ("LuActivity", "Analytics & monitoring", "after every submission"), ("LuQrCode", "Authorised verification", "consent + signed QR")]
    x0, span = 0.45, 12.45 / len(steps)
    for i, (ic, title, sub) in enumerate(steps):
        cx = x0 + i * span + span / 2
        icon(s, ic, cx - 0.3, 2.06, 0.6, bg=ORANGE if i in (0, len(steps) - 1) else NAVY)
        if i < len(steps) - 1:
            arrow(s, cx + 0.4, 2.36, cx + span - 0.4, 2.36, color=BLUE, width=2)
        text(s, cx - span / 2, 2.68, span, 0.28, title, size=12, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        text(s, cx - span / 2, 2.93, span, 0.26, sub, size=11, color=GREY, align=PP_ALIGN.CENTER)
    # Novelty in one line: what is new against the national systems that already exist (detail on slide 6)
    label_box(s, 0.35, 3.24, 12.63, 0.3,
              [[("What is new  ", {"bold": True, "color": ORANGE}),
                ("no national system checks each admission against approved intake. We add that check, on top of APAAR.",
                 {})]],
              fill=PALE_ORANGE, line=ORANGE, size=12, align=PP_ALIGN.CENTER, radius=0.06)

    # How it addresses the problem
    pointer(s, 0.35, 3.64, 6.2, "How it addresses the problem")
    pairs = [("AICTE gets only totals (PS)", "Student-level records on APAAR"),
             ("Excess seat: 5× fee fine each (APH)", "Seat #61 of 60 refused at upload"),
             ("Same student at two colleges", "Cross-college match + contact rings"),
             ("Hard to verify at source (PS)", "Signed QR, checks offline")]
    for i, (prob, sol) in enumerate(pairs):
        y = 4.08 + i * 0.7
        label_box(s, 0.35, y, 2.75, 0.6, [[("✖  ", {"color": RED, "bold": True}), (prob, {})]], fill=PALE_RED,
                  line=RGBColor(0xF0, 0xC4, 0xC0), size=12, align=PP_ALIGN.LEFT)
        arrow(s, 3.15, y + 0.3, 3.65, y + 0.3, color=GREY, width=2)
        label_box(s, 3.7, y, 2.85, 0.6, [[("✔  ", {"color": GREEN, "bold": True}), (sol, {})]], fill=PALE_GREEN,
                  line=RGBColor(0xBF, 0xE3, 0xC9), size=12, align=PP_ALIGN.LEFT)

    # Innovation and uniqueness
    pointer(s, 6.8, 3.64, 6.18, "Innovation and uniqueness of the solution")
    tiles = [("LuUsers2", "Twin-safe matching", "Mohd = Mohammed\nRohit ≠ Mohit"),
             ("LuNetwork", "Ghost-admission map", "1 phone, 8 students,\n2 colleges"),
             ("LuHistory", "Decisions kept", "with a reason;\nnever re-flagged"),
             ("LuStamp", "Signed checkpoints", "history rewrites\ndetected"),
             ("LuScale", "Rules as data", "sourced; unconfirmed\nvalues only warn"),
             ("LuGraduationCap", "APAAR on degrees", "university print\nfeed + QR")]
    tw, th = 1.98, 1.32
    for i, (ic, t1, t2) in enumerate(tiles):
        x, y = 6.8 + (i % 3) * (tw + 0.12), 4.08 + (i // 3) * (th + 0.1)
        box(s, x, y, tw, th, fill=PALE, line=LINE)
        icon(s, ic, x + tw / 2 - 0.25, y + 0.08, 0.5, bg=BLUE)
        text(s, x + 0.05, y + 0.6, tw - 0.1, 0.28, t1, size=12, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        text(s, x + 0.05, y + 0.87, tw - 0.1, 0.42, t2, size=11, color=GREY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)


def slide3(s, live_url, build, video_url=""):
    remove(find(s, "TextBox 8"))
    pointer(s, 0.35, 1.2, 12.63, "Technologies to be used (e.g. programming languages, frameworks, hardware)")
    tech = [("LuServer", "Python · FastAPI", "REST · OpenAPI"), ("LuAtom", "React · TypeScript", "installable PWA"),
            ("LuDatabase", "PostgreSQL 16", "SQLite in the demo"), ("LuBrainCircuit", "scikit-learn", "Isolation Forest"),
            ("LuUsers2", "Name matcher", "twin-safe · Indic"), ("LuLock", "Ed25519 · W3C", "did:web · AES-GCM")]
    cw = 12.63 / len(tech)
    for i, (ic, t1, t2) in enumerate(tech):
        x = 0.35 + i * cw
        box(s, x + 0.04, 1.65, cw - 0.08, 0.78, fill=PALE, line=LINE)
        icon(s, ic, x + 0.12, 1.82, 0.44, bg=NAVY)
        text(s, x + 0.6, 1.72, cw - 0.66, 0.34, t1, size=11, bold=True, color=NAVY)
        text(s, x + 0.6, 2.04, cw - 0.66, 0.3, t2, size=10.5, color=GREY)

    pointer(s, 0.35, 2.6, 12.63, "Methodology and process for implementation (Flow Charts/Images/ working prototype)")
    # Flowchart (left)
    B = dict(fill=PALE, line=LINE, color=NAVY, size=12, bold=True)
    label_box(s, 0.35, 3.1, 1.75, 0.62, "1  Upload\nExcel / CSV / API", **B)
    arrow(s, 2.1, 3.41, 2.35, 3.41)
    label_box(s, 2.35, 3.1, 1.95, 0.62, f"2  Check {RULE_COUNT} rules\n+ APAAR registry", **B)
    arrow(s, 4.3, 3.41, 4.45, 3.41)
    d = box(s, 4.45, 2.98, 1.4, 0.86, fill=PALE_ORANGE, line=ORANGE, shape=MSO_SHAPE.DIAMOND)
    write(d.text_frame, "Errors?", size=11, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)
    d.text_frame.margin_left = d.text_frame.margin_right = Inches(0)
    arrow(s, 5.85, 3.41, 5.95, 3.41, color=GREEN)
    text(s, 5.58, 3.08, 0.36, 0.22, "no", size=11, color=GREEN, bold=True)
    label_box(s, 5.95, 3.1, 1.0, 0.62, "3  Commit", fill=PALE_GREEN, line=RGBColor(0xBF, 0xE3, 0xC9), color=GREEN, size=12, bold=True)
    arrow(s, 5.15, 3.84, 5.15, 4.12, color=RED)
    text(s, 5.2, 3.84, 0.4, 0.22, "yes", size=11, color=RED, bold=True)
    label_box(s, 4.0, 4.12, 2.25, 0.5, "Row returned with reason", fill=PALE_RED, line=RGBColor(0xF0, 0xC4, 0xC0), color=RED, size=12, bold=True)
    arrow(s, 4.0, 4.37, 1.22, 4.37, color=RED, head=False)
    arrow(s, 1.22, 4.37, 1.22, 3.72, color=RED)
    text(s, 1.3, 3.95, 2.0, 0.22, "fix & re-upload", size=11, color=RED)
    arrow(s, 6.45, 3.72, 6.45, 4.85, color=GREY)
    row = [("4  National scan", "dupes · intake · contacts"), ("5  ML risk score", "vs peers, with reasons"),
           ("6  Live dashboard", "decisions remembered")]
    for i, (a, b) in enumerate(row):
        x = 4.95 - i * 2.3
        label_box(s, x, 4.85, 2.0, 0.66, [[(a, {"bold": True})], [(b, {"size": 11, "color": GREY})]], fill=PALE, line=LINE, color=NAVY, size=12)
        if i < len(row) - 1:
            arrow(s, x, 5.18, x - 0.3, 5.18)
    label_box(s, 0.35, 5.85, 3.15, 0.66, [[("7  Student consent (OTP)", {"bold": True})], [("verifier sees only approved fields", {"size": 11, "color": GREY})]],
              fill=PALE_GREEN, line=RGBColor(0xBF, 0xE3, 0xC9), color=GREEN, size=12)
    label_box(s, 3.75, 5.85, 3.2, 0.66, [[("8  Signed QR · degree feed", {"bold": True})], [("APAAR + student ID; forgery fails", {"size": 11, "color": GREY})]],
              fill=PALE_GREEN, line=RGBColor(0xBF, 0xE3, 0xC9), color=GREEN, size=12)
    arrow(s, 1.35, 5.51, 1.35, 5.85)
    arrow(s, 5.35, 5.51, 5.35, 5.85)

    # Working prototype (right)
    h = picture(s, SHOTS / "03_upload_validation.png", 7.25, 3.08, 5.7, crop=(0.178, 0.058, 0, 0.53))
    text(s, 7.25, 3.08 + h + 0.05, 5.7, 0.26, "Working prototype: 77-row upload → 65 accepted, 12 returned with rule codes",
         size=11, color=GREY)
    # What is real and what is simulated, so no integration is mistaken for a live government connection
    y = 3.08 + h + 0.36
    status = [("Built", GREEN, f"upload, {RULE_COUNT} rules, scan, risk score, consent, signed QR, audit"),
              ("Simulated (no public API yet)", ORANGE, "APAAR lookup, DigiLocker, portal sign-in"),
              ("Planned", BLUE, "API Setu / DigiLocker partner access, govt SSO, STQC audit")]
    box(s, 7.22, y, 5.76, 0.9, fill=WHITE, line=LINE)
    for i, (k, col, v) in enumerate(status):
        text(s, 7.3, y + 0.04 + i * 0.28, 5.62, 0.28, [[(f"{k}: ", {"bold": True, "color": col}), (v, {})]],
             size=10.5, color=INK)
    y += 0.98
    box(s, 7.22, y, 5.76, 6.85 - y, fill=PALE_ORANGE, line=ORANGE)
    if live_url:
        qr(s, live_url, 7.32, y + 0.06, 6.85 - y - 0.12, build / "qr_live.png")
        text(s, 7.32 + (6.85 - y - 0.12) + 0.12, y + 0.03, 4.8, 6.85 - y - 0.06,
             [[("Try it live: ", {"bold": True, "color": ORANGE}), (live_url, {"color": BLUE, "link": live_url})],
              ([("Demo video (5 min): ", {"size": 10.5, "color": GREY}), (bare(video_url), {"size": 10.5, "color": BLUE, "link": video_url}),
                (" · 7 roles · synthetic data", {"size": 10.5, "color": GREY})] if video_url else
               [("7 demo roles · synthetic data · first load can take a minute (free server)", {"size": 10.5, "color": GREY})])],
             size=12, anchor=MSO_ANCHOR.MIDDLE)
    else:
        text(s, 7.32, y + 0.05, 5.6, 6.85 - y - 0.1, [[("Run locally: ", {"bold": True, "color": ORANGE}), ("./run.sh", {"color": BLUE})],
             [("Demo users for every role · synthetic data", {"size": 11, "color": GREY})]], size=12)


def slide4(s):
    remove(find(s, "TextBox 8"))
    pointer(s, 0.35, 1.2, 12.63, "Analysis of the feasibility of the idea")
    cur, prev = HELD["current_matcher"], HELD["previous_matcher"]
    of = lambda v: v.replace("/", " of ")  # noqa: E731  "17/17" -> "17 of 17"
    held_twins = of(HELD["current_matcher"]["by_kind"]["twin"].split()[0])
    planted = [BENCH[k].split("/") for k in ("dual_enrolment_recall", "identity_duplicate_recall", "shared_contact_clusters_found")]
    dual, ids, rings = (p[0] for p in planted)
    # Each tile reads as one sentence: the number, then what it counts.
    stats = [("LuUsers2", f"{cur['precision']:.2f} / {cur['recall']:.2f}",
              f"name-match precision / recall on {cur['pairs']} held-out pairs (old {prev['precision']:.2f} / {prev['recall']:.2f})"),
             ("LuUserCheck", of(BENCH["twins_wrongly_flagged"]),
              f"seeded twin pairs wrongly merged (held-out: {held_twins}, sent to an officer)"),
             ("LuShieldCheck", of(BENCH["planted_defect_types_caught"]),
              f"error types caught in the demo upload; 0 false flags on {BENCH['upload_rows_clean']} clean rows"),
             ("LuNetwork", f"{sum(int(p[0]) for p in planted)} of {sum(int(p[1]) for p in planted)}",
              f"planted cases found: {dual} dual enrolments, {ids} ID duplicates, {rings} contact rings"),
             ("LuBrainCircuit", of(BENCH["risk_precision_at_k"]), "highest-risk colleges were the planted anomalies"),
             ("LuFlaskConical", f"{API_TESTS} tests", f"{COVERAGE}% line coverage; security and WCAG checks in CI")]
    sw = 12.63 / len(stats)
    for i, (ic, v, l) in enumerate(stats):
        x = 0.35 + i * sw
        box(s, x + 0.05, 1.62, sw - 0.1, 1.2, fill=PALE, line=LINE)
        icon(s, ic, x + 0.15, 1.7, 0.4, bg=BLUE)
        text(s, x + 0.6, 1.66, sw - 0.68, 0.46, v, size=17 if len(v) > 8 else 20, bold=True, color=NAVY)
        text(s, x + 0.13, 2.13, sw - 0.23, 0.68, l, size=10.5, color=INK, anchor=MSO_ANCHOR.TOP)
    label_box(s, 0.35, 2.87, 12.63, 0.32,
              [[("All synthetic data: ", {"bold": True, "color": RED}),
                ("regression checks, not real-world accuracy (the pilot measures that). An officer decides every flag.", {})]],
              fill=PALE_RED, line=RGBColor(0xF0, 0xC4, 0xC0), size=12, align=PP_ALIGN.CENTER, radius=0.06)

    pointer(s, 0.35, 3.32, 6.0, "Potential challenges and risks")
    pointer(s, 6.98, 3.32, 6.0, "Strategies for overcoming these challenges")
    rows = [("LuPlugZap", "No public APAAR API", "Adapter: mock now → DigiLocker / API Setu"),
            ("LuUserX", "Students without APAAR (opt-out, FN / OCI)", "Provisional ID; never rejected"),
            ("LuLock", "Privacy (DPDP Act; APAAR is voluntary)", "Field-level consent; contacts encrypted; no Aadhaar"),
            ("LuScale", "Wrong flags: twins, spellings, bad data", "Reasons shown; officer decides; decision kept"),
            ("LuCalendarClock", "Rules change; some values unconfirmed", "Rules file per year; unconfirmed values only warn"),
            ("LuSmartphone", "Varied ERP formats, low digital readiness", "Headers auto-mapped; Excel template; Hindi; phone")]
    for i, (ic, ch, st) in enumerate(rows):
        y = 3.74 + i * 0.47
        box(s, 0.35, y, 6.0, 0.41, fill=PALE_RED, line=RGBColor(0xF0, 0xC4, 0xC0))
        icon(s, ic, 0.41, y + 0.04, 0.33, bg=RED)
        text(s, 0.85, y, 5.45, 0.41, ch, size=12, bold=True, color=INK)
        arrow(s, 6.4, y + 0.2, 6.93, y + 0.2, color=GREY, width=2)
        box(s, 6.98, y, 6.0, 0.41, fill=PALE_GREEN, line=RGBColor(0xBF, 0xE3, 0xC9))
        text(s, 7.05, y, 5.88, 0.41, [[("✔  ", {"color": GREEN, "bold": True}), (st, {})]], size=12, color=INK)
    y = 6.6
    box(s, 0.35, y, 12.63, 0.3, fill=PALE_ORANGE, line=ORANGE, radius=0.06)
    text(s, 0.45, y, 12.45, 0.3, [[("Roadmap  ", {"bold": True, "color": ORANGE}),
                                   ("M1–3 pilot: 1 state, 50 colleges, measure hours saved → M4–6 API Setu, SSO, STQC → "
                                    "M7–12 national   ", {}),
                                   ("Cost  ", {"bold": True, "color": ORANGE}),
                                   ("no licence fees", {})]],
         size=11, color=INK)


def slide5(s):
    remove(find(s, "TextBox 8"))
    pointer(s, 0.35, 1.2, 6.6, "Potential impact on the target audience")
    cx, cy, hw = 3.65, 4.0, 2.5
    nodes = [("LuLandmark", "AICTE & regulators", "excess and dual\nenrolment at submission", 0.35, 1.68),
             ("LuBuilding2", "Institutions", "one upload;\nerrors explained", 4.95, 1.68),
             ("LuGraduationCap", "Universities", "APAAR printed\non degrees", 0.35, 3.45),
             ("LuBriefcase", "Employers", "QR check;\nconsent for more", 4.95, 3.45),
             ("LuUsers", "Students", "control who\nsees their record", 2.68, 5.65)]
    for ic, t1, t2, x, y in nodes:
        arrow(s, cx, cy, x + 0.98, y + 0.55, color=LINE, width=2.5, head=False)
    hh = picture(s, SHOTS / "10_aicte_dashboard.png", cx - hw / 2, cy - 0.85, hw, crop=(0.178, 0.058, 0, 0.12))
    text(s, cx - hw / 2, cy - 0.85 + hh + 0.04, hw, 0.24, "One national view", size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    for ic, t1, t2, x, y in nodes:
        box(s, x, y, 1.95, 1.1, fill=WHITE, line=LINE)
        icon(s, ic, x + 0.08, y + 0.1, 0.46, bg=ORANGE if t1 == "Students" else NAVY)
        text(s, x + 0.6, y + 0.08, 1.33, 0.5, t1, size=12, bold=True, color=NAVY)
        text(s, x + 0.08, y + 0.6, 1.85, 0.48, t2, size=11, color=INK, anchor=MSO_ANCHOR.TOP)
    # Before / after, in the free corners of the impact map
    before = ["₹23.1 cr fine: 42 seats × ₹11 lakh × 5 (SC, 2019)",
              "AISHE 2022-23 data out in 2026",
              "APAAR–Aadhaar mismatches (CU Srinagar, 2026)"]
    after = ["Seat #61 of 60 refused at upload",
             "National view after each submission",
             "Mismatch returned with its reason"]
    for x, w, title, items, fill, line, col, mark in (
            (0.35, 2.2, "Today", before, PALE_RED, RGBColor(0xF0, 0xC4, 0xC0), RED, "✖"),
            (4.72, 2.23, "With Vidya Pramaan", after, PALE_GREEN, RGBColor(0xBF, 0xE3, 0xC9), GREEN, "✔")):
        box(s, x, 4.86, w, 1.9, fill=fill, line=line)
        text(s, x + 0.08, 4.9, w - 0.16, 0.3, title, size=12, bold=True, color=col)
        text(s, x + 0.08, 5.18, w - 0.16, 1.56, [[(f"{mark} ", {"color": col, "bold": True}), (t, {})] for t in items],
             size=11, color=INK, anchor=MSO_ANCHOR.TOP, spacing=4)

    pointer(s, 7.2, 1.2, 5.78, "Benefits of the solution (social, economic, environmental, etc.)")
    ben = [("LuHeartHandshake", "Social", "APAAR opt-outs never\nrefused (SC, 2026)", GREEN),
           ("LuIndianRupee", "Economic", "no licence fees; one upload\nserves AICTE + university", BLUE),
           ("LuLeaf", "Environmental", "fewer paper returns\nand letters", GREEN),
           ("LuLandmark", "Governance", "audited evidence\npack for show-cause", NAVY)]
    bw, bh = 2.83, 1.3
    for i, (ic, t1, t2, col) in enumerate(ben):
        x, y = 7.2 + (i % 2) * (bw + 0.12), 1.62 + (i // 2) * (bh + 0.08)
        box(s, x, y, bw, bh, fill=PALE, line=LINE)
        icon(s, ic, x + 0.12, y + 0.1, 0.5, bg=col)
        text(s, x + 0.72, y + 0.1, bw - 0.8, 0.5, t1, size=13, bold=True, color=NAVY)
        text(s, x + 0.12, y + 0.66, bw - 0.2, 0.6, t2, size=12, color=INK, anchor=MSO_ANCHOR.TOP)
    text(s, 7.2, 4.3, 5.78, 0.26, "Expected benefits; the pilot measures them.", size=11,
         color=GREY, align=PP_ALIGN.CENTER)
    y = 4.74
    h1 = picture(s, SHOTS / "25_certificate.png", 7.25, y, 2.5)
    h2 = picture(s, SHOTS / "23_public_qr_verified.png", 10.13, y, 2.82, crop=(0.17, 0.46, 0.51, 0.29))
    text(s, 7.2, y + h1 + 0.05, 2.7, 0.26, "Certificate: APAAR ID + student ID + QR", size=10.5, color=GREY, align=PP_ALIGN.CENTER)
    text(s, 10.1, y + h2 + 0.05, 2.88, 0.26, "QR: on-device check + revocation list", size=10.5, color=GREY, align=PP_ALIGN.CENTER)


def slide6(s, live_url, repo_url, build, video_url=""):
    remove(find(s, "TextBox 8"))
    pointer(s, 0.35, 1.2, 12.63, "Details / Links of the reference and research work")
    text(s, 0.35, 1.6, 7.0, 0.3, "Why not just use the systems that exist?", size=13, bold=True, color=NAVY)
    cols = ["NAD / DigiLocker", "ABC", "AISHE", "Vidya Pramaan"]
    rows = [("Student-level admissions vs approved intake", ["—", "—", "totals, 2–3 yrs late", "✔ per upload"]),
            ("Same student at two colleges; ghost admissions", ["—", "—", "—", "✔"]),
            ("Supernumerary caps by programme and year", ["—", "—", "—", "✔"]),
            ("Consent-based verification for employers", ["✔", "—", "—", "✔ field-level"]),
            ("APAAR + student ID feed for degree printing", ["issuer uploads", "credits only", "—", "✔"])]
    x0, cw0, cw = 0.35, 2.75, 1.06
    widths = [cw0, cw, cw * 0.8, cw * 1.2, cw * 1.05]
    xs = [x0]
    for w in widths[:-1]:
        xs.append(xs[-1] + w)
    y = 1.95
    for j, h in enumerate(["", *cols]):
        if h:
            label_box(s, xs[j] + 0.02, y, widths[j] - 0.04, 0.38, h, fill=ORANGE if h == "Vidya Pramaan" else NAVY,
                      color=WHITE, size=10, bold=True, radius=0.05)
    for i, (cap, vals) in enumerate(rows):
        yy = y + 0.44 + i * 0.47
        box(s, x0, yy, sum(widths), 0.44, fill=PALE if i % 2 == 0 else WHITE, line=LINE, radius=0.04)
        text(s, xs[0] + 0.05, yy, widths[0] - 0.1, 0.44, cap, size=10.5, color=INK)
        for j, v in enumerate(vals, start=1):
            good = v.startswith("✔")
            text(s, xs[j], yy, widths[j], 0.44, v, size=10.5 if len(v) < 12 else 9.5, bold=good,
                 color=GREEN if good else GREY, align=PP_ALIGN.CENTER)
    text(s, 0.35, y + 0.44 + 5 * 0.47 + 0.02, 7.0, 0.62,
         [[("✔ = built and tested (docs/CLAIMS.md). ", {"bold": True, "color": NAVY}),
           ("Reuses APAAR; no new ID.", {})]], size=11, color=GREY, anchor=MSO_ANCHOR.TOP)

    refs = [("AICTE PS SIH26263, SIH 2026", "sih.gov.in"),
            ("AICTE Approval Process Handbook 2024-27 and 2026-27 addendum (FN/OCI/PIO caps)", "aicte.gov.in"),
            ("Supreme Court, FORE School of Management v AICTE (2019): 5× fee per excess student", "sci.gov.in"),
            ("NAD on DigiLocker: consent-based verification (FAQ)", "nad.digilocker.gov.in"),
            ("AISHE 2022-23 & 2023-24 reports (released 2026)", "aishe.gov.in"),
            ("APAAR opt-out: Orissa HC, Rohit Anand Das (Dec 2025); Supreme Court direction (Jul 2026)", "scconline.com · livelaw.in"),
            ("Parliament replies: AICTE enrolment 2025-26 (Lok Sabha); APAAR IDs at 30 Mar 2026 (Rajya Sabha)", "sansad.in"),
            ("Digital Personal Data Protection Act 2023; DPDP Rules 2025", "meity.gov.in"),
            ("W3C Verifiable Credentials and did:web; RFC 8032 (Ed25519)", "w3.org · rfc-editor.org"),
            ("Liu, Ting & Zhou (2008) Isolation Forest; IDinsight on matching Indian names", "doi.org/10.1109/ICDM.2008.17 · idinsight.org")]
    text(s, 7.6, 1.6, 5.4, 0.3, "References", size=13, bold=True, color=NAVY)
    paras = []
    for i, (t, link) in enumerate(refs, 1):
        paras.append([(f"{i}. {t}  ", {"size": 10}), (link, {"size": 9.5, "color": BLUE})])
    text(s, 7.6, 1.92, 5.38, 3.45, paras, size=10, anchor=MSO_ANCHOR.TOP, spacing=2)
    if HAS_QUOTE:
        who = ", ".join(x for x in (FIELD.get("speaker_role"), FIELD.get("institution")) if x)
        text(s, 7.6, 4.88, 5.38, 0.5, [[("From the field:  ", {"bold": True, "color": ORANGE, "size": 10}),
                                       (f"“{FIELD['quote'].strip()}”", {"italic": True, "size": 10}),
                                       (f"  {who}" if who else "", {"size": 9.5, "color": GREY})]],
             size=10, anchor=MSO_ANCHOR.TOP)

    y = 5.45
    box(s, 0.35, y, 12.63, 1.38, fill=PALE_ORANGE, line=ORANGE)
    slots = []
    if video_url:
        slots.append(("Demo video, 5 min", video_url, "qr_video.png", "every role, narrated, captioned"))
    if live_url:
        slots.append(("Live prototype", live_url, "qr_live6.png", "7 demo roles · synthetic data"))
    slots.append(("Source code & tests", repo_url, "qr_repo.png", f"{API_TESTS} tests · {COVERAGE}% coverage · MIT"))
    sw = 12.63 / len(slots)
    for i, (t, url, fn, sub) in enumerate(slots):
        x = 0.35 + i * sw + 0.12
        qr(s, url, x, y + 0.12, 1.14, build / fn)
        text(s, x + 1.24, y + 0.12, sw - 1.4, 1.14,
             [[(t, {"bold": True, "color": ORANGE, "size": 13})],
              [(bare(url), {"color": BLUE, "size": 10.5 if len(slots) > 2 else 12, "link": url})],
              [(sub, {"size": 10.5, "color": GREY})]], size=12, anchor=MSO_ANCHOR.MIDDLE)


NOTES = [
    "Vidya Pramaan: one student, one APAAR ID, one verifiable record. A working prototype for SIH26263, live at the "
    "link and QR on this slide, with seven demo roles and synthetic data only. The screenshots are the prototype: "
    "the AICTE national view, where each college's headline reason is the rule it broke; the same app on a phone; "
    "a probable duplicate under two APAAR IDs with the matcher's reasons; and one phone number shared by eight "
    "students at two colleges. The five-minute narrated demo video shows every role end to end.",
    "29.6 lakh enrolments in AICTE institutions in 2025-26 and 24.25 crore APAAR IDs, yet AICTE today receives totals "
    "only. Excess admission costs five times the fee per student, and it surfaces in annual returns or inspections. Our "
    "flow is the one the problem statement asks for: upload, APAAR check, verification, central record, monitoring "
    "after every submission, consented verification. What is new is the per-student check against approved intake, "
    "which no existing national system makes, plus twin-safe name matching, ghost-admission rings, remembered "
    "decisions, signed audit checkpoints, sourced rules and APAAR on every degree.",
    "Open-source stack, PostgreSQL tested. Every row goes through 31 explained rules; errors come back with reasons; "
    "accepted rows trigger a national scan and risk scoring against peers. The screenshot is the real prototype: "
    "77 rows, 65 accepted, 12 returned. We are explicit about what is simulated: there is no public APAAR API yet, so "
    "the registry lookup is a mock behind an adapter that switches to DigiLocker or API Setu after onboarding.",
    "These are measurements on synthetic data, and we say so. On 49 held-out name pairs: precision 0.79, recall 0.92, "
    "against 0.65 and 0.68 for our earlier matcher. None of the 80 seeded twin pairs is flagged; 2 of 12 held-out twin "
    "pairs are, and go to an officer, who decides. Real-world accuracy is the first thing the pilot measures. "
    f"{API_TESTS} tests on SQLite and PostgreSQL with {COVERAGE}% coverage, plus security and accessibility checks on every push.",
    "Everyone gains from one national view: AICTE and regulators see problems at submission, institutions get "
    "explained errors, universities print APAAR on degrees, employers check a QR, students control access. Every "
    "item on the left is sourced: 42 excess seats times an 11 lakh fee times five is the 23.1 crore fine the Supreme "
    "Court upheld, national statistics arrive years late, and universities report APAAR mismatches this year. The "
    "benefits on the right are expected, not yet measured; the pilot measures them.",
    "We do not replace NAD, ABC or AISHE; we fill the gap none of them covers: student-level reconciliation of "
    "admissions against approved intake. Sources are on the right. The three QR codes open the five-minute demo "
    "video, the live prototype, and the code with its tests.",
]


def add_notes(slides):
    notes = list(NOTES)
    if HAS_QUOTE:
        notes[5] += (f" We spoke to {FIELD.get('people_spoken_to') or 'an'} admission-cell staff; "
                     f"one told us: “{FIELD['quote'].strip()}”")
    for slide, note in zip(slides, notes):
        slide.notes_slide.notes_text_frame.text = note


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--team-id", default="194599")
    ap.add_argument("--team-name", default="AI Riders")
    ap.add_argument("--live-url", default="https://vidya-pramaan-sih26263.onrender.com")
    ap.add_argument("--repo-url", default="https://github.com/Pulkitkh/VidyaPramaan")
    ap.add_argument("--video-url", default="https://youtu.be/eBwySqMs63Q")
    ap.add_argument("--template", default=str(HERE / "template" / "SIH2026-IDEA-Presentation-Format.pptx"))
    ap.add_argument("--out", default=str(HERE / "SIH26263_AI_Riders_Idea.pptx"))
    a = ap.parse_args()
    build = HERE / ".build"
    build.mkdir(exist_ok=True)

    prs = Presentation(a.template)
    # Drop the instructions slide (the template says it may be deleted before upload).
    sld_ids = prs.slides._sldIdLst
    last = sld_ids[-1]
    prs.part.drop_rel(last.get(qn("r:id")))
    sld_ids.remove(last)

    slides = list(prs.slides)
    for s in slides[1:]:
        set_runs(find(s, "Oval"), a.team_name, size=12)
    slide1(slides[0], a.team_id, a.team_name, build, a.live_url, a.video_url)
    slide2(slides[1], a.team_name)
    slide3(slides[2], a.live_url, build, a.video_url)
    slide4(slides[3])
    slide5(slides[4])
    slide6(slides[5], a.live_url, a.repo_url, build, a.video_url)
    add_notes(slides)
    cp = prs.core_properties
    cp.title = "Vidya Pramaan: SIH26263 idea submission"
    cp.author = f"Team {a.team_name}"
    cp.subject = PS_TITLE
    cp.keywords = "SIH 2026, SIH26263, APAAR, AICTE"
    cp.comments = ""
    prs.save(a.out)
    print("wrote", a.out)


if __name__ == "__main__":
    main()
