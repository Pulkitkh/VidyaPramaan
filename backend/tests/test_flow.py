"""End-to-end tests over a small synthetic dataset (isolated temp database)."""
import json
import os

import pytest

from harness import SUMMARY, client, login


@pytest.fixture(scope="module")
def inst():
    return login("inst.admin")


@pytest.fixture(scope="module")
def aicte():
    return login("aicte.admin")


@pytest.fixture(scope="module")
def batch(inst):
    with open(os.environ["VP_SAMPLE_CSV"], "rb") as fh:
        r = client.post("/api/institution/uploads", headers=inst, files={"file": ("cse.csv", fh, "text/csv")})
    assert r.status_code == 200, r.text
    return r.json()


def codes(batch) -> set[str]:
    return {f["code"] for row in batch["rows"] for f in row["findings"]}


def test_upload_flags_every_planted_defect(batch):
    expected = {"APR-001", "APR-002", "APR-003", "APR-004", "APR-005", "APR-007", "DUP-001", "DUP-002", "DUP-003",
                "INT-001", "INT-002", "INT-003", "CON-001", "CON-002", "CON-004", "CON-005", "CON-007"}
    assert expected <= codes(batch)
    assert batch["total_rows"] == 77
    assert batch["error_rows"] == 12


def test_intake_overflow_counts_seats(batch):
    seats = [f["evidence"]["seat_number"] for row in batch["rows"] for f in row["findings"] if f["code"] == "INT-001"]
    assert seats == [61, 62, 63]


def test_rbac_blocks_other_roles(batch):
    student = login("student.ananya")
    assert client.get("/api/dashboard/national", headers=student).status_code == 403
    assert client.post(f"/api/institution/uploads/{batch['id']}/commit", headers=login("aicte.admin")).status_code == 403


def test_commit_never_exceeds_intake(inst, batch):
    r = client.post(f"/api/institution/uploads/{batch['id']}/commit", headers=inst)
    assert r.status_code == 200, r.text
    assert r.json()["created"] == 65
    ov = client.get("/api/institution/overview", headers=inst).json()
    cse = next(p for p in ov["ledger"] if p["code"] == "BTCSE")
    assert cse["within_intake"] == 60 and cse["excess"] == 0
    assert cse["supernumerary"]["TFW"]["used"] == 3
    assert client.post(f"/api/institution/uploads/{batch['id']}/commit", headers=inst).status_code == 409


def test_national_dashboard_and_risk(aicte):
    d = client.get("/api/dashboard/national", headers=aicte).json()
    assert d["kpis"]["institutions"] == 60
    assert d["kpis"]["dual_enrolments"] > 0
    assert d["top_risk"][0]["risk_score"] >= d["top_risk"][-1]["risk_score"]
    planted = {i for ids in SUMMARY["planted"].values() for i in ids}
    top = {i["id"] for i in d["top_risk"][:5]}
    assert len(top & planted) >= 4


def test_consent_flow_discloses_only_scope():
    verifier, student = login("hr.verifier"), login("student.ananya")
    rec = client.get("/api/me/record", headers=student).json()
    apaar = rec["student"]["apaar_id"]
    r = client.post("/api/verify/consent-requests", headers=verifier,
                    json={"apaar_id": apaar, "purpose": "Offer verification", "scope": ["enrolment_status", "programme"]})
    assert r.status_code == 200, r.text
    cid = r.json()["id"]
    assert client.get(f"/api/verify/consent-requests/{cid}/result", headers=verifier).json()["status"] == "PENDING"
    pending = next(c for c in client.get("/api/me/consents", headers=student).json() if c["id"] == cid)
    assert client.post(f"/api/me/consents/{cid}/decision", headers=student,
                       json={"decision": "APPROVE", "otp": "bad"}).status_code == 400
    client.post(f"/api/me/consents/{cid}/decision", headers=student,
                json={"decision": "APPROVE", "otp": pending["demo_otp"]}).raise_for_status()
    res = client.get(f"/api/verify/consent-requests/{cid}/result", headers=verifier).json()
    assert set(res["disclosed"]) == {"enrolment_status", "programme"}
    assert "institution" in res["withheld"]


def test_qr_token_verifies_and_detects_tampering():
    student = login("student.ananya")
    token = client.get("/api/me/record", headers=student).json()["credentials"][0]["token"]
    ok = client.post("/api/verify/token", json={"token": token}).json()
    assert ok["verdict"] == "VERIFIED"
    h, p, s = token.split(".")
    forged = ".".join([h, p[:-2] + ("A" if p[-2] != "A" else "B") + p[-1], s])
    assert client.post("/api/verify/token", json={"token": forged}).json()["verdict"] == "TAMPERED_OR_INVALID"


def test_certificate_pdf():
    student = login("student.ananya")
    token = student["Authorization"].split()[1]
    eid = client.get("/api/me/record", headers=student).json()["enrolments"][0]["id"]
    r = client.get(f"/api/enrolments/{eid}/certificate.pdf", headers=student)
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    # session tokens are never accepted in URLs
    assert client.get(f"/api/enrolments/{eid}/certificate.pdf", params={"token": token}).status_code == 401


def test_audit_chain_detects_tampering(aicte):
    ok = client.get("/api/audit/verify", headers=aicte).json()
    assert ok["valid"] is True and ok["checkpoints_verified"] >= 1
    edit = client.post("/api/audit/simulate", headers=aicte, json={"attack": "edit"}).json()
    assert edit["chain_only"]["valid"] is False and edit["chain_only"]["broken_at"]
    # nothing stored was changed by the simulation
    assert client.get("/api/audit/verify", headers=aicte).json()["valid"] is True


def test_reports(aicte):
    assert client.get("/api/reports/excess-intake.csv", headers=aicte).status_code == 200
    r = client.get("/api/reports/state-summary.xlsx", headers=aicte)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert client.get("/api/reports/excess-intake.csv").status_code == 401


def test_reupload_after_commit_is_update_not_new_admission(inst, batch):
    with open(os.environ["VP_SAMPLE_CSV"], "rb") as fh:
        again = client.post("/api/institution/uploads", headers=inst, files={"file": ("cse.csv", fh, "text/csv")}).json()
    before = {r["row_no"]: {f["code"] for f in r["findings"]} for r in batch["rows"] if r["status"] != "ERROR"}
    for row in again["rows"]:
        codes_now = {f["code"] for f in row["findings"]}
        if row["row_no"] in before:
            assert "UPD-001" in codes_now and row["is_update"]
            # re-uploading must not invent new problems (no seat overflow, no self-match)
            assert codes_now - {"UPD-001"} <= before[row["row_no"]], (row["row_no"], codes_now)
    # the rows that were over intake the first time are still over intake
    assert [f["evidence"]["seat_number"] for r in again["rows"] for f in r["findings"] if f["code"] == "INT-001"] == [61, 62, 63]


def test_reupload_applies_corrections(inst):
    import csv as _csv
    import io as _io
    rows = list(_csv.DictReader(open(os.environ["VP_SAMPLE_CSV"])))
    target = next(r for r in rows if r["apaar_consent"] == "REFUSED")  # no APAAR: matched by roll number
    target["category"] = "SC" if target["category"] != "SC" else "OBC"
    buf = _io.StringIO()
    w = _csv.DictWriter(buf, fieldnames=list(rows[0]))
    w.writeheader()
    w.writerow(target)
    b = client.post("/api/institution/uploads", headers=inst,
                    files={"file": ("fix.csv", buf.getvalue().encode(), "text/csv")}).json()
    assert [f["code"] for f in b["rows"][0]["findings"]] == ["APR-005", "UPD-001"]
    res = client.post(f"/api/institution/uploads/{b['id']}/commit", headers=inst).json()
    assert res["created"] == 0 and res["updated"] == 1
    students = client.get(f"/api/institution/students?q={target['roll_no']}", headers=inst).json()
    assert students[0]["student"]["category"] == target["category"]


def test_excel_upload_matches_csv(inst):
    import csv as _csv
    import io as _io
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for r in _csv.reader(open(os.environ["VP_SAMPLE_CSV"])):
        ws.append([int(v) if v.isdigit() and len(v) == 12 else v for v in r])  # Excel turns IDs into numbers
    buf = _io.BytesIO()
    wb.save(buf)
    x = client.post("/api/institution/uploads", headers=inst, files={"file": ("cse.xlsx", buf.getvalue(), "application/vnd.ms-excel")})
    assert x.status_code == 200, x.text
    assert x.json()["total_rows"] == 77
    t = client.get("/api/institution/template.xlsx")
    assert t.status_code == 200 and t.content[:2] == b"PK"


def test_demo_reset_restores_empty_programme(aicte, inst):
    r = client.post("/api/demo/reset", headers=aicte)
    assert r.status_code == 200 and r.json()["removed_enrolments"] == 65
    cse = next(p for p in client.get("/api/institution/overview", headers=inst).json()["ledger"] if p["code"] == "BTCSE")
    assert cse["within_intake"] == 0
    assert client.get("/api/audit/verify", headers=aicte).json()["valid"] is True
    assert client.post("/api/demo/reset", headers=inst).status_code == 403


def test_expired_token_rejected(monkeypatch):
    from app import auth
    monkeypatch.setattr(auth, "TOKEN_TTL_SECONDS", -1)
    stale = auth.issue_token(auth.User("aicte.admin", "AICTE_ADMIN", "x"))
    r = client.get("/api/dashboard/national", headers={"Authorization": f"Bearer {stale}"})
    assert r.status_code == 401 and "expired" in r.json()["detail"]


def test_exports_neutralise_formulas():
    from app.main import _cell
    assert _cell("=HYPERLINK(\"x\")") == "'=HYPERLINK(\"x\")"
    assert _cell("Pune") == "Pune" and _cell(5) == 5


# --------------------------------------------------------------------------- claims made in the idea deck
@pytest.fixture(scope="module")
def committed(inst):
    """A fresh upload + commit of the demo file (runs after the demo-reset test)."""
    with open(os.environ["VP_SAMPLE_CSV"], "rb") as fh:
        b = client.post("/api/institution/uploads", headers=inst, files={"file": ("cse.csv", fh, "text/csv")}).json()
    r = client.post(f"/api/institution/uploads/{b['id']}/commit", headers=inst)
    assert r.status_code == 200, r.text
    return b


def test_rule_catalogue_has_31_explained_rules(committed):
    from app.validation import RULES
    assert len(RULES) == 31
    for row in committed["rows"]:
        for f in row["findings"]:
            assert f["code"] in RULES and f["message"] and isinstance(f["evidence"], dict) and f["rule_version"]


def test_json_api_validate_and_submit(inst):
    import csv as _csv
    rows = list(_csv.DictReader(open(os.environ["VP_SAMPLE_CSV"])))[:3]
    v = client.post("/api/v1/students:validate", headers=inst, json=rows)
    assert v.status_code == 200 and len(v.json()) == 3
    s = client.post("/api/v1/students:submit", headers=inst, json=rows)
    assert s.status_code == 200 and s.json()["source"] == "API"
    spec = client.get("/openapi.json").json()
    assert "/api/v1/students:submit" in spec["paths"] and "/api/v1/students:validate" in spec["paths"]


def test_provisional_ids_for_students_without_apaar(inst, committed):
    rows = client.get("/api/institution/students?limit=200", headers=inst).json()
    no_apaar = [r for r in rows if not r["student"]["apaar_id"]]
    assert no_apaar and all(r["student"]["provisional_id"].startswith("PRV-") for r in no_apaar)


def test_officer_decision_is_audited(aicte, inst):
    f = client.get("/api/findings?limit=1", headers=aicte).json()["items"][0]
    assert client.post(f"/api/findings/{f['id']}/decision", headers=inst, json={"decision": "RESOLVED"}).status_code == 403
    r = client.post(f"/api/findings/{f['id']}/decision", headers=aicte,
                    json={"decision": "RESOLVED", "note": "verified with institution"})
    assert r.status_code == 200 and r.json()["status"] == "RESOLVED"
    log = client.get("/api/audit?action=FINDING_RESOLVED", headers=aicte).json()["items"]
    assert any(e["entity_id"] == str(f["id"]) and e["details"]["note"] == "verified with institution" for e in log)


def test_auditor_is_read_only(aicte):
    aud = login("auditor")
    assert client.get("/api/dashboard/national", headers=aud).status_code == 200
    f = client.get("/api/findings?limit=1", headers=aud).json()["items"][0]
    assert client.post(f"/api/findings/{f['id']}/decision", headers=aud, json={"decision": "DISMISSED"}).status_code == 403
    assert client.post("/api/demo/reset", headers=aud).status_code == 403
    assert client.post("/api/scan", headers=aud).status_code == 403
    # reading certificate data never issues a credential on the auditor's behalf
    from app.db import SessionLocal
    from app.models import AuditEvent, Credential, Enrolment
    with SessionLocal() as db:
        issued = {c.enrolment_id for c in db.query(Credential)}
        eid = next(e.id for e in db.query(Enrolment).limit(500) if e.id not in issued)
        before = db.query(AuditEvent).count()
    assert client.get(f"/api/v1/certificate-data/{eid}", headers=aud).status_code == 404
    with SessionLocal() as db:
        assert db.query(AuditEvent).count() == before


def test_consent_revocation_and_access_log():
    verifier, student = login("hr.verifier"), login("student.ananya")
    apaar = client.get("/api/me/record", headers=student).json()["student"]["apaar_id"]
    cid = client.post("/api/verify/consent-requests", headers=verifier,
                      json={"apaar_id": apaar, "purpose": "Revocation check", "scope": ["programme"]}).json()["id"]
    otp = next(c for c in client.get("/api/me/consents", headers=student).json() if c["id"] == cid)["demo_otp"]
    client.post(f"/api/me/consents/{cid}/decision", headers=student, json={"decision": "APPROVE", "otp": otp}).raise_for_status()
    assert client.get(f"/api/verify/consent-requests/{cid}/result", headers=verifier).json()["disclosed"]
    log = client.get("/api/me/record", headers=student).json()["access_log"]
    assert any(l["purpose"] == "Revocation check" and l["basis"] == "CONSENT" for l in log)
    client.post(f"/api/me/consents/{cid}/revoke", headers=student).raise_for_status()
    after = client.get(f"/api/verify/consent-requests/{cid}/result", headers=verifier).json()
    assert after["status"] == "REVOKED" and after["disclosed"] is None


def test_apaar_ids_are_masked_in_aicte_views(aicte):
    import re
    items = client.get("/api/findings?limit=200", headers=aicte).json()["items"]
    blob = json.dumps(items)
    assert not re.search(r"(?<!\d)\d{12}(?!\d)", blob), "full APAAR IDs must not appear in the review queue"
    detail = client.get(f"/api/findings/{items[0]['id']}", headers=aicte).json()
    if "record" in detail:
        assert detail["record"]["student"]["apaar_id"].startswith("XXXX-") or detail["record"]["student"]["apaar_id"] == "—"


def test_risk_scores_come_with_plain_reasons(aicte):
    top = client.get("/api/dashboard/national", headers=aicte).json()["top_risk"][0]
    assert top["risk_score"] >= 70 and top["reasons"]
    rules = [r for r in top["reasons"] if r.get("rule")]
    stats = [r for r in top["reasons"] if not r.get("rule")]
    assert all("peer median" in r["text"] for r in stats)
    # A broken cap is the headline reason, ahead of any peer comparison.
    assert rules and top["reasons"][:len(rules)] == rules
    assert all(r["rule"] in ("INT-001", "INT-002") and f"(rule {r['rule']})" in r["text"] for r in rules)


def test_excel_template_has_dropdowns_and_report_has_diversity(aicte):
    import io as _io
    from openpyxl import load_workbook
    wb = load_workbook(_io.BytesIO(client.get("/api/institution/template.xlsx").content))
    assert len(wb.active.data_validations.dataValidation) >= 9
    rep = load_workbook(_io.BytesIO(client.get("/api/reports/state-summary.xlsx", headers=aicte).content))
    assert {"State summary", "Findings by rule", "Diversity", "Ruleset"} <= set(rep.sheetnames)


def test_spa_route_cannot_read_files_outside_the_build():
    from app.main import DIST
    if not DIST.exists():
        pytest.skip("frontend not built")
    # frontend/dist -> ../../backend/app/config.py always exists, so an escape would be visible.
    for prefix in ("%2e%2e/%2e%2e/", "..%2f..%2f", "%2e%2e%2f%2e%2e%2f"):
        r = client.get(f"/{prefix}backend/app/config.py")
        assert b"RULES_FILE" not in r.content, prefix
