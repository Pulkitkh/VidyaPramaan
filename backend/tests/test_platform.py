"""Tests for the review workflow, identity matching, lifecycle, roles, integration and security hardening."""
import csv
import io
import os
from datetime import date, timedelta

import pytest

from harness import SUMMARY, client, login

from app.credentials import verify_compact  # noqa: E402  (after harness sets the environment)


@pytest.fixture(scope="module")
def aicte():
    return login("aicte.admin")


@pytest.fixture(scope="module")
def inst():
    return login("inst.admin")


@pytest.fixture(scope="module", autouse=True)
def fresh_demo(aicte):
    assert client.post("/api/demo/reset", headers=aicte).status_code == 200


def sample_rows() -> list[dict]:
    return list(csv.DictReader(open(os.environ["VP_SAMPLE_CSV"], encoding="utf-8")))


def as_csv(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=list(rows[0]))
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue().encode()


def validate(inst, rows) -> list[dict]:
    r = client.post("/api/v1/students:validate", headers=inst, json=rows)
    assert r.status_code == 200, r.text
    return r.json()


def codes(row) -> list[str]:
    return [f["code"] for f in row["findings"]]


def clean_row(**kw) -> dict:
    """A valid, previously unseen row for the demo institution (opt-out student, so no registry needed)."""
    r = {"apaar_id": "", "apaar_consent": "REFUSED", "roll_no": "T-0001", "name": "Kavitha Raman", "dob": "2008-02-11",
         "gender": "F", "category": "OBC", "father_name": "Raman Iyer", "mother_name": "", "guardian_name": "",
         "programme_code": "BTECE", "year_of_study": "1", "admission_type": "REGULAR", "quota": "STATE_CAP",
         "nationality": "Indian", "fn_status": "N", "oci_status": "N", "oci_card_no": "", "passport_no": "",
         "nri_status": "N", "nri_country": "", "pwd": "N", "admission_date": "2026-08-03", "mobile": "",
         "email": ""}
    r.update(kw)
    return r


# --------------------------------------------------------------------------- identity matching
def test_matcher_keeps_twins_apart_and_joins_spelling_variants():
    from app.matching import Person, compare_names, same_person, tokens
    d = date(2007, 2, 27)
    twin = same_person(Person("Rohit Sharma", d, "M", "Vinod Sharma"), Person("Mohit Sharma", d, "M", "Vinod Sharma"))
    assert not twin.match
    same = same_person(Person("Mohammed Irfan", d, "M", "Shri Abdul Rashid"), Person("Mohd Irfan", d, "M", "Abdul Rashid"))
    assert same.match and same.reasons
    assert compare_names("Venkata Sai Krishna", "V S Krishna").match
    assert compare_names("Lakshmi Narayanan", "Laxmi Narayan").match
    assert tokens("राहुल शर्मा") == ["राहुल", "शर्मा"]       # vowel signs are kept
    assert tokens("கார்த்திக்") == ["கார்த்திக்"]
    swapped = same_person(Person("Kiran Reddy", date(2006, 4, 7), "M", "N Reddy"),
                          Person("Kiran Reddy", date(2006, 7, 4), "M", "N Reddy"))
    assert swapped.match and "swapped" in swapped.reasons[0]


def test_identity_matching_beats_previous_matcher_on_held_out_pairs():
    from scripts.identity_eval import SETS, evaluate, load
    res = evaluate(load(SETS["held_out"]))
    cur, prev = res["current_matcher"], res["previous_matcher"]
    assert cur["recall"] > prev["recall"] and cur["precision"] > prev["precision"]
    assert cur["recall"] >= 0.85 and cur["precision"] >= 0.75


def test_seeded_twins_are_not_flagged_as_duplicates(aicte):
    from app.db import SessionLocal
    from app.models import Enrolment, Finding
    with SessionLocal() as db:
        twins = {e.id for e in db.query(Enrolment) if 600 <= int(e.roll_no[-3:]) < 700}
        flagged = {f.enrolment_id for f in db.query(Finding).filter(Finding.rule_code == "DUP-003")}
    assert len(twins) == 80 and not (twins & flagged)


def test_within_file_duplicate_without_apaar_is_caught(inst):
    a = clean_row(roll_no="T-0101", name="Rahul Kumar Sharma", father_name="Suresh Sharma", dob="2008-05-02")
    b = clean_row(roll_no="T-0102", name="Rahul K. Sharma", father_name="Shri Suresh Sharma", dob="2008-05-02")
    res = validate(inst, [a, b])
    assert "DUP-003" not in codes(res[0]) and "DUP-003" in codes(res[1])


def test_shared_contact_clusters_are_found_with_members(aicte):
    d = client.get("/api/dashboard/national", headers=aicte).json()
    assert d["kpis"]["shared_contact_clusters"] == len(SUMMARY["contact_clusters"])
    f = client.get("/api/findings?rule=DUP-004&limit=1", headers=aicte).json()["items"][0]
    detail = client.get(f"/api/findings/{f['id']}", headers=aicte).json()
    assert len(detail["cluster"]) >= 4 and detail["evidence"]["institutions"] >= 2
    assert "XXXXXX" in detail["evidence"]["contact"] or "***@" in detail["evidence"]["contact"]


def test_kpis_count_cases_not_rows(aicte):
    d = client.get("/api/dashboard/national", headers=aicte).json()
    assert d["kpis"]["dual_enrolments"] == d["cases_by_rule"]["DUP-002"] < d["findings_by_rule"]["DUP-002"]


# --------------------------------------------------------------------------- review workflow
def test_decisions_need_a_reason_and_survive_the_next_scan(aicte):
    f = client.get("/api/findings?rule=DUP-003&limit=1", headers=aicte).json()["items"][0]
    assert client.post(f"/api/findings/{f['id']}/decision", headers=aicte,
                       json={"decision": "DISMISSED", "note": ""}).status_code == 400
    client.post(f"/api/findings/{f['id']}/decision", headers=aicte,
                json={"decision": "DISMISSED", "note": "Twins, verified birth certificates"}).raise_for_status()
    open_before = {x["id"] for x in client.get("/api/findings?rule=DUP-003&limit=200", headers=aicte).json()["items"]}
    scan = client.post("/api/scan", headers=aicte).json()
    assert scan["decisions_kept"] >= 1
    after = client.get(f"/api/findings/{f['id']}", headers=aicte).json()
    assert after["status"] == "DISMISSED" and after["resolution_note"].startswith("Twins")
    open_after = {x["id"] for x in client.get("/api/findings?rule=DUP-003&limit=200", headers=aicte).json()["items"]}
    assert open_after == open_before  # same ids, nothing re-opened


# --------------------------------------------------------------------------- validation details
def test_consent_is_never_assumed(inst):
    from app.db import SessionLocal
    from app.models import MockApaarRecord
    with SessionLocal() as db:
        reg = db.query(MockApaarRecord).filter(MockApaarRecord.status == "ACTIVE").all()[-1]
    row = clean_row(apaar_id=reg.apaar_id, apaar_consent="", name=reg.name, dob=reg.dob.isoformat(), gender=reg.gender)
    res = validate(inst, [row])[0]
    assert "APR-008" in codes(res) and res["data"]["apaar_consent"] == "NOT_RECORDED"


def test_mother_or_guardian_is_enough(inst):
    rows = [clean_row(roll_no="T-0201", father_name="", mother_name="Lata Raman"),
            clean_row(roll_no="T-0202", name="Arif Khan", father_name="", guardian_name="Salim Khan (uncle)"),
            clean_row(roll_no="T-0203", name="No Parent", father_name="")]
    res = validate(inst, rows)
    assert "SCH-001" not in codes(res[0]) and res[0]["status"] != "ERROR"
    assert "SCH-002" in codes(res[1])  # "(uncle)" is not part of a name
    assert "SCH-001" in codes(res[2])


def test_values_that_cannot_be_read_are_rejected(inst):
    future = (date.today() + timedelta(days=30)).isoformat()
    rows = [clean_row(roll_no="T-0301", fn_status="maybe"), clean_row(roll_no="T-0302", year_of_study="abc"),
            clean_row(roll_no="T-0303", admission_date=future), clean_row(roll_no="T-0304", name="<script>x</script>"),
            clean_row(roll_no="T-0305", mobile="12345"), clean_row(roll_no="T-0306", email="not-an-email")]
    res = validate(inst, rows)
    assert [r["status"] for r in res] == ["ERROR"] * 6
    assert "CON-009" in codes(res[2])


def test_erp_headers_are_mapped_automatically(inst):
    erp = [{"Student Name": "Meera Pillai", "APAAR No": "", "Consent": "No", "Roll Number": "T-0401",
            "Date of Birth": "11-02-2008", "Sex": "Female", "Social Category": "General", "Mother's Name": "Asha Pillai",
            "Course Code": "BTECE", "Entry Type": "Regular", "Admission Quota": "State Cap",
            "Date of Admission": "03-08-2026", "Mobile No": "+91 98765 43210", "Hostel Block": "C"}]
    b = client.post("/api/institution/uploads", headers=inst,
                    files={"file": ("erp_export.csv", as_csv(erp), "text/csv")}).json()
    assert b["column_map"]["renamed"]["Student Name"] == "name"
    assert b["column_map"]["ignored"] == ["Hostel Block"]
    row = b["rows"][0]
    assert row["status"] == "OK", row["findings"]
    assert row["data"]["gender"] == "F" and row["data"]["mobile"] == "9876543210"


def test_near_cap_note_and_supernumerary_caps_by_level(aicte):
    from app import config
    assert config.sn_cap_fraction("FN_OCI_PIO", "UG") == 0.15 and config.sn_cap_fraction("FN_OCI_PIO", "PG") == 0.25
    rules = client.get("/api/rules").json()
    assert rules["ruleset"]["academic_year"] == "2026-27"
    assert rules["ruleset"]["supernumerary"]["FN_OCI_PIO"]["verified"] is True
    ov = client.get("/api/institutions/1", headers=aicte).json()
    mba = next(p for p in ov["ledger"] if p["level"] == "PG")
    assert mba["supernumerary"]["FN_OCI_PIO"]["cap"] == int(mba["approved_intake"] * 0.25)


def test_continuing_students_do_not_take_this_years_seats(inst):
    ok = clean_row(roll_no="T-0501", admission_year="2025", year_of_study="2", admission_date="2025-08-01")
    wrong_year = clean_row(roll_no="T-0502", name="Anil Joshi", admission_year="2025", year_of_study="1",
                           admission_date="2025-08-01")
    res = validate(inst, [ok, wrong_year])
    assert res[0]["status"] == "OK" and not {"INT-001", "INT-003"} & set(codes(res[0]))
    assert "CON-010" in codes(res[1])


# --------------------------------------------------------------------------- commit and lifecycle
def test_commit_revalidates_a_stale_batch(inst):
    data = open(os.environ["VP_SAMPLE_CSV"], "rb").read()
    first = client.post("/api/institution/uploads", headers=inst, files={"file": ("a.csv", data, "text/csv")}).json()
    second = client.post("/api/institution/uploads", headers=inst, files={"file": ("b.csv", data, "text/csv")}).json()
    assert client.post(f"/api/institution/uploads/{first['id']}/commit", headers=inst).json()["created"] == 65
    res = client.post(f"/api/institution/uploads/{second['id']}/commit", headers=inst).json()
    # the second batch was validated before the first was committed: it must become updates, not new seats
    assert res["created"] == 0 and res["changed_since_validation"]
    cse = next(p for p in client.get("/api/institution/overview", headers=inst).json()["ledger"] if p["code"] == "BTCSE")
    assert cse["within_intake"] == 60


def test_correction_preview_supersedes_the_old_credential(inst):
    student = login("student.ananya")
    rec = client.get("/api/me/record", headers=student).json()
    e = rec["enrolments"][0]
    old_token = rec["credentials"][0]["token"]
    rows = client.get(f"/api/institution/students?q={e['roll_no']}", headers=inst).json()
    s = rows[0]["student"]
    row = clean_row(apaar_id=s["apaar_id"], apaar_consent="GIVEN", roll_no=e["roll_no"], name="Ananya S Kulkarni",
                    dob=s["dob"], gender="F", category=s["category"], father_name=s["father_name"] or "",
                    mother_name=s["mother_name"] or "", programme_code=e["programme_code"],
                    admission_date=e["admission_date"])
    b = client.post("/api/institution/uploads", headers=inst, files={"file": ("fix.csv", as_csv([row]), "text/csv")}).json()
    upd = next(f for f in b["rows"][0]["findings"] if f["code"] == "UPD-001")
    assert {"field": "name", "old": "Ananya Kulkarni", "new": "Ananya S Kulkarni"} in upd["evidence"]["changes"]
    assert "UPD-002" in codes(b["rows"][0])
    assert client.post(f"/api/institution/uploads/{b['id']}/commit", headers=inst).json()["updated"] == 1
    assert client.post("/api/verify/token", json={"token": old_token}).json()["verdict"] == "SUPERSEDED"


def test_exit_closes_dual_enrolment_and_revokes_credential(aicte, inst):
    dup = next(f for f in client.get("/api/findings?rule=DUP-002&limit=200", headers=inst).json()["items"])
    eid = dup["enrolment_id"]
    cd = client.get(f"/api/v1/certificate-data/{eid}", headers=inst).json()
    bad = client.post(f"/api/institution/enrolments/{eid}/exit", headers=inst,
                      json={"status": "TRANSFERRED_OUT", "date": "2030-01-01", "reason": "NOC issued"})
    assert bad.status_code == 400
    r = client.post(f"/api/institution/enrolments/{eid}/exit", headers=inst,
                    json={"status": "TRANSFERRED_OUT", "date": date.today().isoformat(), "reason": "Migrated with NOC"})
    assert r.status_code == 200 and r.json()["credentials_revoked"] == 1
    still = [f for f in client.get("/api/findings?rule=DUP-002&limit=200", headers=inst).json()["items"]
             if f["enrolment_id"] == eid]
    assert not still
    status = client.get(f"/api/credentials/{cd['credential_id']}/status").json()
    assert status["status"] == "REVOKED" and status["enrolment_status"] == "TRANSFERRED_OUT"
    # Offline verifiers learn about the revocation from the signed list they cache while online.
    listed = verify_compact(client.get("/api/credentials/status-list").json()["token"])
    assert listed[0] and cd["credential_id"] in listed[1]["revoked"]


def test_issuer_is_a_resolvable_did_web_with_the_signing_key():
    doc = client.get("/.well-known/did.json")
    assert doc.status_code == 200 and doc.headers["content-type"].startswith("application/did+json")
    d = doc.json()
    assert d["id"].startswith("did:web:") and "example" not in d["id"]
    jwk = d["verificationMethod"][0]["publicKeyJwk"]
    assert jwk["x"] == client.get("/api/issuer/jwk").json()["x"]
    assert d["assertionMethod"] == [f"{d['id']}#key-1"]
    lst = client.get("/api/credentials/status-list").json()
    ok, claims = verify_compact(lst["token"])
    assert ok and claims["iss"] == d["id"] and claims["typ"] == "revocation-list" and claims["exp"] > claims["iat"]
    # An edited list fails its signature, so it cannot hide a revocation.
    h, p, sg = lst["token"].split(".")
    assert not verify_compact(f"{h}.{p[:-2]}AA.{sg}")[0]


# --------------------------------------------------------------------------- roles
def test_university_gets_degree_print_feed_for_affiliated_colleges_only(aicte):
    uni = login("university.registrar")
    ov = client.get("/api/university/overview", headers=uni).json()
    ids = {i["id"] for i in ov["institutions"]}
    assert 1 in ids and len(ids) > 1
    feed = client.get("/api/university/degree-feed?institution_id=1", headers=uni).json()
    row = next(r for r in feed["items"] if r["apaar_id"])
    assert len(row["apaar_id"]) == 12 and row["student_id"]
    other = next(i for i in client.get("/api/institutions?limit=200", headers=aicte).json()["items"] if i["id"] not in ids)
    assert client.get(f"/api/university/degree-feed?institution_id={other['id']}", headers=uni).status_code == 403
    csv_r = client.get("/api/university/degree-feed.csv", headers=uni)
    assert csv_r.status_code == 200 and "APAAR ID" in csv_r.text.splitlines()[0]
    assert client.get(f"/api/v1/certificate-data/{row['enrolment_id']}", headers=uni).status_code == 200
    log = client.get("/api/audit?action=DEGREE_FEED_EXPORTED", headers=aicte).json()["items"]
    assert log and log[0]["role"] == "UNIVERSITY"


def test_regulator_sees_aggregates_for_its_state_only():
    reg = login("dte.regulator")
    assert client.get("/api/dashboard/national", headers=reg).status_code == 200
    items = client.get("/api/institutions?limit=200", headers=reg).json()["items"]
    assert items and {i["state"] for i in items} == {"Maharashtra"}
    assert client.get("/api/findings", headers=reg).status_code == 403
    assert client.get("/api/audit", headers=reg).status_code == 403


# --------------------------------------------------------------------------- integration API
def test_api_keys_and_idempotent_submit(inst):
    k = client.post("/api/institution/api-keys", headers=inst, json={"label": "ERP sync"}).json()
    assert k["key"].startswith("vpk_")
    key = {"X-API-Key": k["key"]}
    rows = [clean_row(roll_no="T-0601", name="Sneha Kulkarni", programme_code="MBA", dob="2003-04-04")]
    first = client.post("/api/v1/students:submit", headers={**key, "Idempotency-Key": "sync-42"}, json=rows).json()
    assert first["created"] == 1
    again = client.post("/api/v1/students:submit", headers={**key, "Idempotency-Key": "sync-42"}, json=rows).json()
    assert again["replayed"] is True and again["id"] == first["id"]
    assert not any(x.get("key") for x in client.get("/api/institution/api-keys", headers=inst).json())
    client.delete(f"/api/institution/api-keys/{k['id']}", headers=inst).raise_for_status()
    assert client.post("/api/v1/students:validate", headers=key, json=rows).status_code == 401


# --------------------------------------------------------------------------- consent and verification hardening
def test_otp_is_hashed_and_locks_after_five_wrong_tries():
    verifier, student = login("hr.verifier"), login("student.ananya")
    apaar = client.get("/api/me/record", headers=student).json()["student"]["apaar_id"]
    cid = client.post("/api/verify/consent-requests", headers=verifier,
                      json={"apaar_id": apaar, "purpose": "Background verification", "scope": ["programme"]}).json()["id"]
    from app.db import SessionLocal
    from app.models import ConsentRequest
    with SessionLocal() as db:
        c = db.get(ConsentRequest, cid)
        assert len(c.otp_hash) == 64 and c.demo_otp not in c.otp_hash
    for _ in range(5):
        r = client.post(f"/api/me/consents/{cid}/decision", headers=student, json={"decision": "APPROVE", "otp": "000001"})
        assert r.status_code == 400
    mine = next(c for c in client.get("/api/me/consents", headers=student).json() if c["id"] == cid)
    assert mine["status"] == "LOCKED"
    assert client.post(f"/api/me/consents/{cid}/revoke", headers=student).status_code == 409


def test_consent_request_does_not_reveal_whether_an_apaar_exists():
    verifier = login("hr.verifier")
    r = client.post("/api/verify/consent-requests", headers=verifier,
                    json={"apaar_id": "999999999999", "purpose": "Background verification", "scope": ["programme"]})
    assert r.status_code == 200 and r.json()["status"] == "PENDING"


def test_public_qr_scans_do_not_flood_the_access_log():
    student = login("student.ananya")
    eid = client.get("/api/me/record", headers=student).json()["enrolments"][0]["id"]
    client.get(f"/api/v1/certificate-data/{eid}", headers=student).raise_for_status()  # re-issued after correction
    rec = client.get("/api/me/record", headers=student).json()
    token = next(c["token"] for c in rec["credentials"] if c["status"] == "ACTIVE")
    assert client.post("/api/verify/token", json={"token": token}).json()["verdict"] == "VERIFIED"
    before = len(client.get("/api/me/record", headers=student).json()["access_log"])
    for _ in range(5):
        client.post("/api/verify/token", json={"token": token})
    after = len(client.get("/api/me/record", headers=student).json()["access_log"])
    assert after - before <= 1


def test_event_stream_needs_a_short_lived_ticket(aicte):
    token = aicte["Authorization"].split()[1]
    assert client.get("/api/events", params={"ticket": token}).status_code == 401
    t = client.post("/api/events/ticket", headers=aicte).json()["ticket"]
    from app.auth import parse_ticket
    assert parse_ticket(t, "events").role == "AICTE_ADMIN"


# --------------------------------------------------------------------------- audit
def test_signed_checkpoints_catch_rewrites_and_truncation(aicte):
    rewrite = client.post("/api/audit/simulate", headers=aicte, json={"attack": "rewrite"}).json()
    assert rewrite["chain_only"]["valid"] is True            # a plain hash chain is fooled
    assert rewrite["with_checkpoints"]["valid"] is False
    assert rewrite["with_checkpoints"]["reason"] == "history_rewritten"
    trunc = client.post("/api/audit/simulate", headers=aicte, json={"attack": "truncate"}).json()
    assert trunc["chain_only"]["valid"] is True and trunc["with_checkpoints"]["reason"] == "events_missing"
    pub = client.get("/api/audit/checkpoints").json()
    assert pub["items"] and pub["items"][0]["signature"].count(".") == 2


def test_issuer_key_rotation_resigns_everything():
    from app.db import SessionLocal
    from app.keys import ensure_current_key
    from app.models import Setting
    with SessionLocal() as db:
        db.get(Setting, "issuer_kid").value = "previous-key"
        db.commit()
        res = ensure_current_key(db)
    assert res["credentials_resigned"] >= 1 and res["checkpoints_resigned"] >= 1
    aicte = login("aicte.admin")
    assert client.get("/api/audit/verify", headers=aicte).json()["valid"] is True
    log = client.get("/api/audit?action=ISSUER_KEY_ROTATED", headers=aicte).json()["items"]
    assert log and log[0]["details"]["old_key"] == "previous-key"


# --------------------------------------------------------------------------- deployment safety
def test_security_headers_are_set():
    r = client.get("/api/meta")
    h = r.headers
    assert "frame-ancestors 'none'" in h["content-security-policy"]
    assert h["x-content-type-options"] == "nosniff" and h["cache-control"] == "no-store"
    assert "commit" in r.json()  # the deploy workflow waits for this to match the pushed commit


def test_demo_reset_needs_the_pin_when_one_is_set(monkeypatch, aicte):
    from app import config
    monkeypatch.setattr(config, "DEMO_PIN", "4821")
    assert client.post("/api/demo/reset", headers=aicte, json={"pin": "0000"}).status_code == 403
    assert client.post("/api/demo/reset", headers=aicte, json={"pin": "4821"}).status_code == 200


def test_demo_conveniences_are_off_outside_demo_mode(monkeypatch):
    from app import config
    monkeypatch.setattr(config, "DEMO_MODE", False)
    assert client.post("/api/auth/demo-login", json={"username": "aicte.admin"}).status_code == 404


def test_evidence_pack_for_show_cause(aicte):
    top = client.get("/api/dashboard/national", headers=aicte).json()["top_risk"][0]
    r = client.get(f"/api/reports/institutions/{top['id']}/evidence.xlsx", headers=aicte)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    assert {"Summary", "Seat ledger", "Open findings"} <= set(wb.sheetnames)


def test_demo_reset_undoes_portal_corrections(aicte):
    client.post("/api/demo/reset", headers=aicte).raise_for_status()  # (an earlier test already reset once)
    student = login("student.ananya")
    rec = client.get("/api/me/record", headers=student).json()
    assert rec["student"]["name"] == "Ananya Kulkarni"
    assert all(c["status"] == "ACTIVE" for c in rec["credentials"])
    assert client.get("/api/audit/verify", headers=aicte).json()["valid"] is True


def test_sign_in_works_while_a_slow_scan_holds_the_write_lock(monkeypatch):
    """On a small instance the background scan is slow; users must still be able to sign in and work."""
    import threading
    import time
    from app import main, services
    real = services.score_institutions

    def slow(db):
        from sqlalchemy import update
        from app.models import Institution
        out = real(db)
        # Write, so SQLite's write lock is held, as a slow scoring phase on a small CPU would hold it.
        db.execute(update(Institution).where(Institution.id == 1).values(quality_score=Institution.quality_score))
        time.sleep(9)     # well past SQLite's default 5-second lock wait
        return out

    monkeypatch.setattr(services, "score_institutions", slow)
    t = threading.Thread(target=main.refresh_in_background)
    t.start()
    time.sleep(2)
    r = client.post("/api/auth/demo-login", json={"username": "aicte.admin"})
    t.join()
    assert r.status_code == 200, r.text


def test_unverified_caps_send_rows_for_review_instead_of_rejecting(monkeypatch, inst):
    from app import config

    def int002(rows):
        out = validate(inst, rows)
        items = out["rows"] if isinstance(out, dict) else out
        return [f for row in items for f in row["findings"] if f["code"] == "INT-002"]

    verified = int002(sample_rows())
    assert verified and all(f["severity"] == "ERROR" and f["evidence"]["cap_verified"] for f in verified)
    monkeypatch.setitem(config.SUPERNUMERARY["TFW"], "verified", False)
    unverified = int002(sample_rows())
    assert unverified and all(f["severity"] == "WARNING" and "not yet confirmed" in f["message"] for f in unverified)


# --------------------------------------------------------------------------- spreadsheet edge cases
def test_spreadsheet_exports_from_real_offices_are_read_correctly():
    from openpyxl import Workbook

    from app.validation import _date, normalise, parse_csv, parse_xlsx
    rows = sample_rows()[:3]
    header = list(rows[0])
    # Semicolon-separated, Windows-1252, with blank lines: Excel's "CSV" save on many office PCs.
    text = ";".join(header) + "\r\n\r\n" + "\r\n".join(";".join(r[h] for h in header) for r in rows) + "\r\n;;;\r\n"
    parsed = parse_csv(text.encode("cp1252"))
    assert len(parsed) == 3 and parsed[0]["roll_no"] == rows[0]["roll_no"]
    # Workbook whose first sheet is instructions and whose data sheet has a title row above the header.
    wb = Workbook()
    wb.active.title = "Read me"
    wb.active.append(["Fill one row per student"])
    ws = wb.create_sheet("Data")
    ws.append(["Demo Institute of Technology: admissions 2026-27"])
    ws.append(header)
    for r in rows:
        ws.append([int(r[h]) if h == "apaar_id" and r[h].isdigit() else r[h] for h in header])
    buf = io.BytesIO()
    wb.save(buf)
    got = parse_xlsx(buf.getvalue())
    assert len(got) == 3 and got[0]["apaar_id"] == rows[0]["apaar_id"]
    # Excel date serials, month names, scientific-notation IDs, invisible characters and stray spaces.
    assert _date("39448") == date(2008, 1, 1) and _date("21-Jun-2008") == date(2008, 6, 21)
    r, problems = normalise({"apaar_id": "8.26869E+11", "name": "  Kavitha​   Raman ", "mobile": "9.18028E+09"})
    assert r["apaar_id"] == "" and r["name"] == "Kavitha Raman"
    assert sorted(p["evidence"]["field"] for p in problems if "scientific" in p["message"]) == ["apaar_id", "mobile"]
    assert normalise({"apaar_id": "826868856122.0"})[0]["apaar_id"] == "826868856122"


def test_uploading_the_same_file_twice_creates_nothing_new(inst):
    rows = [clean_row(roll_no="T-TWICE-1", name="Meenakshi Sundaram", father_name="Sundaram Pillai")]
    first = client.post("/api/institution/uploads", headers=inst, files={"file": ("a.csv", as_csv(rows), "text/csv")}).json()
    assert client.post(f"/api/institution/uploads/{first['id']}/commit", headers=inst).json()["created"] == 1
    again = client.post("/api/institution/uploads", headers=inst, files={"file": ("a.csv", as_csv(rows), "text/csv")}).json()
    assert again["rows"][0]["is_update"]
    res = client.post(f"/api/institution/uploads/{again['id']}/commit", headers=inst).json()
    assert res["created"] == 0


def test_audit_chain_cannot_fork_or_record_without_its_lock(monkeypatch):
    from sqlalchemy.exc import IntegrityError

    from app import audit
    from app.db import SessionLocal
    from app.models import AuditEvent
    with SessionLocal() as db:
        head = db.scalars(audit.select(AuditEvent).order_by(AuditEvent.id.desc()).limit(1)).first()
        db.add(AuditEvent(ts=head.ts, actor="x", role="x", action="FORK", entity="x", entity_id="1", details={},
                          prev_hash=head.prev_hash, hash="f" * 64))
        with pytest.raises(IntegrityError):
            db.flush()
        db.rollback()
    # If the lock cannot be taken, the change is refused (503), never written outside the chain order.
    class Stuck:
        def acquire(self, timeout=None):
            return False
    monkeypatch.setattr(audit, "_LOCK", Stuck())
    with SessionLocal() as db, pytest.raises(audit.AuditUnavailable):
        audit.record(db, "x", "x", "TEST", "x", 1)


def test_contact_details_and_document_numbers_are_encrypted_at_rest():
    from sqlalchemy import text

    from app.db import SessionLocal
    from app.models import Student
    with SessionLocal() as db:
        s = db.scalars(Student.__table__.select().where(Student.mobile.is_not(None)).limit(1)).first()
        raw = db.execute(text("SELECT mobile, email FROM student WHERE id = :i"), {"i": s}).one()
        student = db.get(Student, s)
        assert raw.mobile.startswith("enc1:") and student.mobile not in raw.mobile
        assert student.mobile.isdigit() and len(student.mobile) == 10
        raw_pp = db.execute(text("SELECT passport_no FROM student WHERE passport_no IS NOT NULL LIMIT 1")).scalar()
        assert raw_pp.startswith("enc1:")
