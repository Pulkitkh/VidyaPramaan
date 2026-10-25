"""Explainable validation engine for student-level submissions.

Each check emits a finding with a stable rule code, severity, plain-language
message, evidence and the ruleset version, so institutions know exactly what to
fix and AICTE can audit why a record was flagged. Flags are review items, not
verdicts of fraud.

Severity semantics:
  ERROR   row is rejected and must be corrected and re-submitted
  WARNING row is accepted but opens a review item for AICTE / the institution
  INFO    row is accepted; the note is recorded for transparency
"""
import csv
import io
import math
import re
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta

from sqlalchemy import select
from sqlalchemy.orm import Session

from . import config
from .apaar_registry import get_registry, is_valid_format, mask
from .matching import Person, compare_names, norm_name, same_person  # noqa: F401  (norm_name re-exported)
from .models import Enrolment, Institution, Programme, Student

RULES: dict[str, tuple[str, str]] = {
    "SCH-001": ("ERROR", "Required field missing"),
    "SCH-002": ("ERROR", "Invalid field value"),
    "SCH-003": ("ERROR", "Invalid date"),
    "PRG-001": ("ERROR", "Programme not approved for this institution"),
    "APR-001": ("ERROR", "APAAR ID missing"),
    "APR-002": ("ERROR", "APAAR ID format invalid"),
    "APR-003": ("ERROR", "APAAR ID not found in registry"),
    "APR-004": ("WARNING", "Details do not match APAAR registry"),
    "APR-005": ("INFO", "Student opted out of APAAR, provisional ID issued"),
    "APR-006": ("ERROR", "APAAR ID is deactivated"),
    "APR-007": ("INFO", "APAAR not applicable (foreign national / OCI), provisional ID issued"),
    "APR-008": ("WARNING", "APAAR consent status not recorded"),
    "DUP-001": ("ERROR", "Duplicate record within this upload"),
    "DUP-002": ("WARNING", "Already actively enrolled at another institution"),
    "DUP-003": ("WARNING", "Possible same person under a different APAAR ID"),
    "DUP-004": ("WARNING", "Same mobile or e-mail used by many students across institutions"),
    "UPD-001": ("INFO", "Already submitted; details will be updated"),
    "UPD-002": ("WARNING", "Identity details of an existing record would change"),
    "INT-001": ("ERROR", "Admission exceeds approved intake"),
    "INT-002": ("ERROR", "Supernumerary quota exceeded"),
    "INT-003": ("INFO", "Programme is nearly full"),
    "CON-001": ("ERROR", "Foreign National without passport number"),
    "CON-002": ("ERROR", "OCI status without OCI card number"),
    "CON-003": ("WARNING", "Nationality conflicts with FN/OCI status"),
    "CON-004": ("WARNING", "NRI category without country of residence"),
    "CON-005": ("ERROR", "Lateral entry must be in year 2"),
    "CON-006": ("WARNING", "Admission date outside admission window"),
    "CON-007": ("WARNING", "Age implausible for programme level"),
    "CON-008": ("ERROR", "Quota inconsistent with nationality status"),
    "CON-009": ("ERROR", "Admission date is in the future"),
    "CON-010": ("WARNING", "Year of study does not fit the admission year"),
}

COLUMNS = [
    "apaar_id", "apaar_consent", "roll_no", "name", "dob", "gender", "category", "father_name", "mother_name",
    "guardian_name", "programme_code", "admission_year", "year_of_study", "admission_type", "quota",
    "nationality", "fn_status", "oci_status", "oci_card_no", "passport_no", "nri_status", "nri_country", "pwd",
    "admission_date", "mobile", "email",
]
REQUIRED = ["roll_no", "name", "dob", "gender", "category", "programme_code", "admission_type", "quota",
            "admission_date"]
PARENT_FIELDS = ("father_name", "mother_name", "guardian_name")
ENUMS = {
    "gender": {"M", "F", "T"},
    "category": {"GEN", "OBC", "SC", "ST", "EWS"},
    "admission_type": {"REGULAR", "LATERAL"},
    "quota": {"STATE_CAP", "MANAGEMENT", "NRI", "TFW", "FN_OCI_PIO", "EWS_SN", "LATERAL"},
    "apaar_consent": {"GIVEN", "REFUSED"},
}
YES_NO = ("fn_status", "oci_status", "nri_status", "pwd")
SUPERNUMERARY_BY_QUOTA = {"TFW": "TFW", "FN_OCI_PIO": "FN_OCI_PIO", "EWS_SN": "EWS", "LATERAL": "LATERAL"}
MAX_ROWS = 5000
MAX_XLSX_UNCOMPRESSED = 60_000_000

# Institutions export from many ERPs; common header spellings are recognised automatically.
HEADER_SYNONYMS = {
    "apaar_id": ["apaar", "apaar_no", "apaar_number", "apaar_id_no", "abc_id", "apaar_abc_id", "apaar_abc"],
    "apaar_consent": ["consent", "apaar_consent_status", "consent_status"],
    "roll_no": ["roll", "roll_number", "enrolment_no", "enrollment_no", "enrolment_number", "enrollment_number",
                "registration_no", "reg_no", "student_id", "university_roll_no"],
    "name": ["student_name", "full_name", "name_of_student", "name_of_the_student", "candidate_name"],
    "dob": ["date_of_birth", "birth_date", "d_o_b", "birthdate"],
    "gender": ["sex"],
    "category": ["social_category", "caste_category", "reservation_category"],
    "father_name": ["fathers_name", "father", "name_of_father"],
    "mother_name": ["mothers_name", "mother", "name_of_mother"],
    "guardian_name": ["guardians_name", "guardian", "name_of_guardian"],
    "programme_code": ["program_code", "course_code", "branch_code", "programme", "course"],
    "admission_year": ["year_of_admission", "batch", "admission_batch"],
    "year_of_study": ["current_year", "study_year", "year"],
    "admission_type": ["entry_type", "type_of_admission"],
    "quota": ["admission_quota", "seat_quota"],
    "nationality": ["country_of_citizenship", "citizenship"],
    "fn_status": ["foreign_national", "is_foreign_national", "fn"],
    "oci_status": ["oci", "is_oci"],
    "oci_card_no": ["oci_card_number", "oci_no"],
    "passport_no": ["passport", "passport_number"],
    "nri_status": ["nri", "is_nri"],
    "nri_country": ["country_of_residence", "nri_country_of_residence"],
    "pwd": ["pwd_status", "divyang", "person_with_disability", "disability"],
    "admission_date": ["date_of_admission", "admitted_on", "doa"],
    "mobile": ["mobile_no", "mobile_number", "phone", "phone_no", "contact_no", "contact_number"],
    "email": ["email_id", "e_mail", "mail_id", "email_address"],
}
VALUE_SYNONYMS = {
    "gender": {"MALE": "M", "FEMALE": "F", "TRANSGENDER": "T", "OTHER": "T", "BOY": "M", "GIRL": "F"},
    "category": {"GENERAL": "GEN", "OPEN": "GEN", "UR": "GEN", "OBC-NCL": "OBC", "OBCNCL": "OBC", "OBC_NCL": "OBC"},
    "apaar_consent": {"YES": "GIVEN", "Y": "GIVEN", "NO": "REFUSED", "N": "REFUSED", "DECLINED": "REFUSED",
                      "OPTED_OUT": "REFUSED"},
}
_TRUE, _FALSE = {"Y", "YES", "TRUE", "1"}, {"N", "NO", "FALSE", "0", ""}
NAME_RE = re.compile(r"^[^\W\d_](?:[^\W\d_]|[\s.'\-])*$")
MOBILE_RE = re.compile(r"^[6-9]\d{9}$")
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$")


def finding(code: str, message: str, **evidence) -> dict:
    severity, title = RULES[code]
    return {"code": code, "severity": severity, "title": title, "message": message,
            "evidence": evidence, "rule_version": config.RULESET_VERSION}


def _key(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", (h or "").strip().lower().replace("'", "").replace("’", "")).strip("_")


_SYN = {_key(s): c for c, syns in HEADER_SYNONYMS.items() for s in syns}


def map_columns(raw_rows: list[dict]) -> tuple[list[dict], dict, list[str]]:
    """Rename recognised header variants to template columns. Returns rows, the mapping used, unknown headers."""
    headers = list(raw_rows[0]) if raw_rows else []
    mapping, unknown = {}, []
    for h in headers:
        k = _key(h)
        target = k if k in COLUMNS else _SYN.get(k)
        if target and target not in mapping.values():
            mapping[h] = target
        elif h:
            unknown.append(h)
    rows = [{mapping[h]: v for h, v in r.items() if h in mapping} for r in raw_rows]
    renamed = {h: c for h, c in mapping.items() if _key(h) != c}
    return rows, renamed, unknown


EXCEL_EPOCH = date(1899, 12, 30)
SCIENTIFIC_RE = re.compile(r"^\d(?:\.\d+)?E\+?\d+$", re.I)
_INVISIBLE = dict.fromkeys(map(ord, "\u200b\u200c\u200d\u2060\ufeff"))


def _date(v) -> date | None:
    """ISO dates, or day-first Indian formats (DD-MM-YYYY, DD/MM/YYYY, 21-Jun-2008). Month-first is never
    assumed. A bare 5-digit number is an Excel date serial (the cell lost its date format on export)."""
    v = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%b-%Y", "%d %b %Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(v, fmt).date()
        except ValueError:
            pass
    if v.isdigit() and 20000 <= int(v) <= 60000:  # 1954-10-03 .. 2064-04-08
        return EXCEL_EPOCH + timedelta(days=int(v))
    return None


def _decode(content: bytes) -> str:
    """UTF-8 (with or without BOM), else Windows-1252: what Excel's plain "CSV" save produces on Indian PCs."""
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("cp1252", errors="replace")


def parse_csv(content: bytes) -> list[dict]:
    """Comma, semicolon, tab or pipe separated; blank lines and fully empty rows are skipped."""
    text = _decode(content)
    first = text.split("\n", 1)[0]
    delimiter = max(",;\t|", key=first.count) if first else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    out = []
    for row in reader:
        clean = {(k or "").strip(): (v or "").strip() for k, v in row.items() if k is not None}
        if any(clean.values()):
            out.append(clean)
    return out


def _header_score(values) -> int:
    keys = [_key(str(v or "")) for v in values]
    return sum(1 for k in keys if k in COLUMNS or k in _SYN)


def parse_xlsx(content: bytes) -> list[dict]:
    """The first worksheet that looks like student data: its header is the first row (within the top ten)
    with at least three recognised column names, so title rows and instruction sheets are skipped.
    Excel dates and numbers are converted to text."""
    from openpyxl import load_workbook
    with zipfile.ZipFile(io.BytesIO(content)) as z:  # refuse zip bombs before openpyxl inflates them
        if sum(i.file_size for i in z.infolist()) > MAX_XLSX_UNCOMPRESSED:
            raise ValueError("Workbook is too large when uncompressed")
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        header = None
        for _ in range(10):
            values = next(it, None)
            if values is None:
                break
            if _header_score(values) >= 3:
                header = [str(h or "").strip() for h in values]
                break
        if header is not None:
            break
    else:
        return []
    if header is None:
        return []
    out = []
    for values in it:
        if values is None or all(v in (None, "") for v in values):
            continue
        row = {}
        for h, v in zip(header, values):
            if isinstance(v, datetime):
                v = v.date().isoformat()
            elif isinstance(v, date):
                v = v.isoformat()
            elif isinstance(v, bool):
                v = "Y" if v else "N"
            elif isinstance(v, float) and v.is_integer():
                v = int(v)
            row[h] = "" if v is None else str(v).strip()
        out.append(row)
        if len(out) > MAX_ROWS:
            break
    return out


def normalise(raw: dict) -> tuple[dict, list[dict]]:
    """Clean one row into template form. Values that cannot be read become SCH-002 findings."""
    problems: list[dict] = []
    r = {c: " ".join(str(raw.get(c) or "").translate(_INVISIBLE).split()) for c in COLUMNS}
    for c in ("apaar_id", "mobile"):
        if SCIENTIFIC_RE.match(r[c]):
            problems.append(finding("SCH-002", f"{c} '{r[c]}' was shortened by Excel to scientific notation and its "
                                    "digits are lost; format the column as Text and export again", field=c, value=r[c]))
            r[c] = ""
    r["apaar_id"] = re.sub(r"[\s-]", "", r["apaar_id"])
    if r["apaar_id"].endswith(".0") and r["apaar_id"][:-2].isdigit():
        r["apaar_id"] = r["apaar_id"][:-2]
    for c in ("gender", "category", "admission_type", "quota", "apaar_consent", "programme_code"):
        v = r[c].upper().replace(" ", "_") if c in ("admission_type", "quota") else r[c].upper()
        r[c] = VALUE_SYNONYMS.get(c, {}).get(v, v)
    r["apaar_consent"] = r["apaar_consent"] or "NOT_RECORDED"
    r["nationality"] = r["nationality"] or "Indian"
    for c in YES_NO:
        v = r[c].upper()
        if v in _TRUE:
            r[c] = True
        elif v in _FALSE:
            r[c] = False
        else:
            problems.append(finding("SCH-002", f"{c} '{r[c]}' must be Y or N", field=c, value=r[c]))
            r[c] = False
    for c, lo, hi, default in (("year_of_study", 1, 6, 1), ("admission_year", 2000, config.ADMISSION_YEAR,
                                                               config.ADMISSION_YEAR)):
        v = r[c]
        if not v:
            r[c] = default
        elif v.isdigit() and lo <= int(v) <= hi:
            r[c] = int(v)
        else:
            problems.append(finding("SCH-002", f"{c} '{v}' must be a whole number from {lo} to {hi}",
                                    field=c, value=v))
            r[c] = default
    mobile = re.sub(r"[\s-]", "", r["mobile"])
    mobile = mobile[3:] if mobile.startswith("+91") else mobile[2:] if len(mobile) == 12 and mobile.startswith("91") else mobile
    r["mobile"] = mobile
    if mobile and not MOBILE_RE.match(mobile):
        problems.append(finding("SCH-002", "mobile must be a 10-digit Indian mobile number", field="mobile"))
    if r["email"] and not EMAIL_RE.match(r["email"]):
        problems.append(finding("SCH-002", f"email '{r['email']}' is not a valid address", field="email"))
    r["email"] = r["email"].lower()
    for c in ("name", *PARENT_FIELDS):
        if r[c] and not NAME_RE.match(r[c]):
            problems.append(finding("SCH-002", f"{c} contains characters that are not part of a name",
                                    field=c, value=r[c][:40]))
    return r, problems


def age_on(dob: date, on: date) -> int:
    return on.year - dob.year - ((on.month, on.day) < (dob.month, dob.day))


def person_of(r: dict, dob: date | None) -> Person:
    return Person(r["name"], dob, r["gender"], r["father_name"] or None, r["mother_name"] or None,
                  r["guardian_name"] or None)


def person_of_student(s: Student) -> Person:
    return Person(s.name, s.dob, s.gender, s.father_name, s.mother_name, s.guardian_name)


def swapped(d: date) -> date | None:
    try:
        return date(d.year, d.day, d.month) if d.day <= 12 and d.day != d.month else None
    except ValueError:
        return None


class IntakeLedger:
    """Running seat count per programme for the current admission year: committed plus accepted so far."""

    def __init__(self, db: Session, programmes: dict[str, Programme]):
        self.progs = {p.id: p for p in programmes.values()}
        self.intake = {p.id: p.approved_intake for p in programmes.values()}
        self.within: dict[int, int] = defaultdict(int)
        self.supernumerary: dict[tuple[int, str], int] = defaultdict(int)
        self.overflow: dict[tuple[int, str | None], int] = defaultdict(int)  # rejected attempts
        self.near_noted: set[int] = set()
        if not programmes:
            return
        rows = db.execute(
            select(Enrolment.programme_id, Enrolment.seat_type, Enrolment.supernumerary_type)
            .where(Enrolment.programme_id.in_(list(self.intake)), Enrolment.status == "ACTIVE",
                   Enrolment.admission_year == config.ADMISSION_YEAR)
        ).all()
        for pid, seat, sn in rows:
            if seat == "WITHIN_INTAKE":
                self.within[pid] += 1
            else:
                self.supernumerary[(pid, sn)] += 1
        for pid, n in self.within.items():
            if n >= math.ceil(self.intake[pid] * config.NEAR_CAP_WARNING):
                self.near_noted.add(pid)

    def cap(self, pid: int, sn_type: str) -> int:
        return math.floor(self.intake[pid] * config.sn_cap_fraction(sn_type, self.progs[pid].level))


def _changes(e: Enrolment, r: dict, dob: date | None, adm: date | None) -> list[dict]:
    """Field-by-field preview of what an update would change on the stored record."""
    s = e.student
    out = []
    for f in ("name", "gender", "category", *PARENT_FIELDS, "nationality", "fn_status", "oci_status", "oci_card_no",
              "passport_no", "nri_status", "nri_country", "pwd", "mobile", "email"):
        new = r[f] if isinstance(r[f], bool) else (r[f] or None)
        if f in ("name", "gender", "category", "nationality") and not new:
            continue
        old = getattr(s, f)
        if old != new and not (old in (None, "") and new in (None, "")):
            out.append({"field": f, "old": old, "new": new})
    if dob and dob != s.dob:
        out.append({"field": "dob", "old": s.dob.isoformat(), "new": dob.isoformat()})
    if adm and adm != e.admission_date:
        out.append({"field": "admission_date", "old": e.admission_date.isoformat(), "new": adm.isoformat()})
    return out


IDENTITY_FIELDS = {"name", "dob", "gender"}


def validate_rows(db: Session, institution: Institution, raw_rows: list[dict]) -> list[dict]:
    registry = get_registry(db)
    programmes = {p.code: p for p in db.scalars(
        select(Programme).where(Programme.institution_id == institution.id,
                                Programme.academic_year == config.CURRENT_ACADEMIC_YEAR))}
    ledger = IntakeLedger(db, programmes)
    parsed = [normalise(r) for r in raw_rows]
    rows = [p[0] for p in parsed]

    apaars = {r["apaar_id"] for r in rows if r["apaar_id"]}
    existing: dict[str, list[Enrolment]] = defaultdict(list)
    if apaars:
        for e in db.scalars(select(Enrolment).join(Student)
                            .where(Student.apaar_id.in_(apaars), Enrolment.status == "ACTIVE")):
            existing[e.student.apaar_id].append(e)

    # This institution's committed enrolments by roll number, so re-uploads of students without an
    # APAAR ID (opt-out, FN/OCI) are recognised as updates rather than new admissions.
    own_by_roll = {e.roll_no: e for e in db.scalars(
        select(Enrolment).where(Enrolment.institution_id == institution.id, Enrolment.status == "ACTIVE"))}

    dobs = {d for d in (_date(r["dob"]) for r in rows) if d}
    dobs |= {s for s in (swapped(d) for d in list(dobs)) if s}
    by_dob: dict[date, list[Student]] = defaultdict(list)
    if dobs:
        for s in db.scalars(select(Student).where(Student.dob.in_(dobs))):
            by_dob[s.dob].append(s)

    win_start, win_end = (date.fromisoformat(d) for d in config.ADMISSION_WINDOW)
    today = date.today()
    seen_apaar: dict[str, int] = {}
    seen_roll: dict[str, int] = {}
    seen_people: dict[date, list[tuple[int, Person]]] = defaultdict(list)
    results = []

    for idx, (r, problems) in enumerate(parsed):
        row_no = idx + 2  # spreadsheet line number (line 1 is the header)
        f: list[dict] = list(problems)
        missing = [c for c in REQUIRED if not r[c]]
        if missing:
            f.append(finding("SCH-001", f"Missing: {', '.join(missing)}", fields=missing))
        if not any(r[c] for c in PARENT_FIELDS):
            f.append(finding("SCH-001", "Give at least one of father's, mother's or guardian's name",
                             fields=list(PARENT_FIELDS)))
        for col, allowed in ENUMS.items():
            if r[col] and r[col] not in allowed and not (col == "apaar_consent" and r[col] == "NOT_RECORDED"):
                f.append(finding("SCH-002", f"{col} '{r[col]}' is not one of {sorted(allowed)}",
                                 field=col, value=r[col]))
        dob, adm = _date(r["dob"]), _date(r["admission_date"])
        for col, val in (("dob", dob), ("admission_date", adm)):
            if r[col] and val is None:
                f.append(finding("SCH-003", f"{col} '{r[col]}' is not a valid date (use YYYY-MM-DD or DD-MM-YYYY)",
                                 field=col, value=r[col]))
        if dob and dob > today:
            f.append(finding("SCH-003", f"Date of birth {dob} is in the future", field="dob"))
        if adm and adm > today + timedelta(days=1):
            f.append(finding("CON-009", f"Admission date {adm} has not happened yet", admission_date=str(adm)))

        prog = programmes.get(r["programme_code"])
        if r["programme_code"] and prog is None:
            f.append(finding("PRG-001", f"Programme '{r['programme_code']}' is not approved for "
                             f"{config.CURRENT_ACADEMIC_YEAR}", approved=sorted(programmes)))
        current_cohort = r["admission_year"] == config.ADMISSION_YEAR

        # --- APAAR identity ---------------------------------------------------
        apaar = r["apaar_id"]
        is_update = False
        update_of: Enrolment | None = None
        if not apaar:
            if r["fn_status"] or r["oci_status"]:
                f.append(finding("APR-007", "APAAR needs Aadhaar-based KYC, so it is optional for "
                                 "FN/OCI students. A provisional AICTE ID will be issued."))
            elif r["apaar_consent"] == "REFUSED":
                f.append(finding("APR-005", "No APAAR ID because the student declined consent. A "
                                 "provisional AICTE ID will be issued; this is not a defect."))
            else:
                f.append(finding("APR-001", "APAAR ID is blank and no consent refusal is recorded"))
        elif not is_valid_format(apaar):
            f.append(finding("APR-002", f"'{r['apaar_id']}' is not a 12-digit APAAR ID", value=apaar))
        else:
            if r["apaar_consent"] == "NOT_RECORDED":
                f.append(finding("APR-008", "apaar_consent is blank. Record GIVEN (or REFUSED) explicitly; "
                                 "consent is never assumed"))
            reg = registry.lookup(apaar)
            if not reg.found:
                f.append(finding("APR-003", f"{mask(apaar)} does not exist in the APAAR registry",
                                 apaar=mask(apaar), source=reg.source))
            elif reg.status != "ACTIVE":
                f.append(finding("APR-006", f"{mask(apaar)} is {reg.status} in the registry"))
            else:
                nm = compare_names(r["name"], reg.name)
                dob_ok = dob == reg.dob
                if not nm.match or not dob_ok:
                    f.append(finding("APR-004", f"Submitted '{r['name']}' / {r['dob']} vs registry "
                                     f"'{reg.name}' / {reg.dob}", name_match=nm.match,
                                     name_reason=nm.reasons[0] if nm.reasons else None,
                                     submitted_name=r["name"], registry_name=reg.name, dob_match=dob_ok))
            if apaar in seen_apaar:
                f.append(finding("DUP-001", f"Same APAAR ID already appears on row {seen_apaar[apaar]}",
                                 first_row=seen_apaar[apaar]))
            else:
                seen_apaar[apaar] = row_no
            for e in existing.get(apaar, []):
                if e.institution_id == institution.id:
                    if prog and e.programme_id == prog.id:
                        is_update, update_of = True, e
                else:
                    f.append(finding(
                        "DUP-002", f"Active enrolment at {e.institution.name} ({e.institution.aicte_pid}), "
                        f"{e.programme.name}, admitted {e.admission_date}",
                        other_institution=e.institution.name, other_pid=e.institution.aicte_pid,
                        other_institution_id=e.institution_id, other_programme=e.programme.name,
                        other_admission_date=str(e.admission_date)))
        if r["roll_no"]:
            if r["roll_no"] in seen_roll:
                f.append(finding("DUP-001", f"Roll number already used on row {seen_roll[r['roll_no']]}",
                                 first_row=seen_roll[r["roll_no"]]))
            else:
                seen_roll[r["roll_no"]] = row_no
            own = own_by_roll.get(r["roll_no"])
            if own and not is_update and prog and own.programme_id == prog.id:
                if (own.student.apaar_id or "") == apaar:
                    is_update, update_of = True, own
                else:
                    holder = mask(own.student.apaar_id) if own.student.apaar_id else own.student.provisional_id
                    f.append(finding("DUP-001", f"Roll number {r['roll_no']} already belongs to "
                                     f"{own.student.name} ({holder})", existing_enrolment=own.id))
        changes: list[dict] = []
        if update_of is not None:
            changes = _changes(update_of, r, dob, adm)
            ignored = "; quota / seat changes need AICTE approval and are not applied" \
                if r["quota"] and r["quota"] != update_of.quota else ""
            summary = ", ".join(c["field"] for c in changes) or "no field changes"
            f.append(finding("UPD-001", f"Already submitted as {update_of.roll_no}; will update: {summary}{ignored}",
                             enrolment_id=update_of.id, changes=changes))
            touched = {c["field"] for c in changes} & IDENTITY_FIELDS
            if touched:
                impact = ["the issued credential and QR are replaced by a new one",
                          "the record is re-checked against the APAAR registry",
                          "an AICTE review item records the old and new values"]
                f.append(finding("UPD-002", f"Changing {', '.join(sorted(touched))} of an existing student "
                                 "needs a reason on file", fields=sorted(touched), impact=impact,
                                 changes=[c for c in changes if c["field"] in touched]))

        # --- probabilistic identity match (same person, different APAAR) ------
        if dob and r["name"] and r["gender"]:
            me = person_of(r, dob)
            for d in {dob, swapped(dob)} - {None}:
                hit = None
                for s in by_dob.get(d, []):
                    if (update_of is not None and s.id == update_of.student_id) or (apaar and s.apaar_id == apaar):
                        continue
                    m = same_person(me, person_of_student(s))
                    if m.match:
                        hit = (s, m)
                        break
                if hit:
                    s, m = hit
                    other = db.scalars(select(Enrolment).where(Enrolment.student_id == s.id)).first()
                    where = other.institution.name if other else "unknown institution"
                    f.append(finding(
                        "DUP-003", f"Matches {s.name} (APAAR {mask(s.apaar_id)}) at {where}: {'; '.join(m.reasons[:3])}",
                        score=m.score, reasons=m.reasons, matched_apaar=mask(s.apaar_id), matched_name=s.name,
                        other_institution=other.institution.name if other else None,
                        other_institution_id=other.institution_id if other else None))
                    break
            for prev_row, p in seen_people[dob]:
                m = same_person(me, p)
                if m.match and (not apaar or apaar != rows[prev_row - 2]["apaar_id"]):
                    f.append(finding("DUP-003", f"Probably the same person as row {prev_row} ({p.name}): "
                                     f"{'; '.join(m.reasons[:3])}", score=m.score, reasons=m.reasons,
                                     first_row=prev_row))
                    break
            seen_people[dob].append((row_no, me))

        # --- consistency --------------------------------------------------------
        foreign = r["nationality"].lower() not in {"indian", "india"}
        if r["fn_status"] and not r["passport_no"]:
            f.append(finding("CON-001", "FN status is Yes but passport number is blank"))
        if r["oci_status"] and not r["oci_card_no"]:
            f.append(finding("CON-002", "OCI status is Yes but OCI card number is blank"))
        if (r["fn_status"] and not foreign) or (foreign and not (r["fn_status"] or r["oci_status"])):
            f.append(finding("CON-003", f"Nationality '{r['nationality']}' with FN={r['fn_status']}, "
                             f"OCI={r['oci_status']}"))
        if (r["nri_status"] or r["quota"] == "NRI") and not r["nri_country"]:
            f.append(finding("CON-004", "NRI category needs country of residence"))
        if r["quota"] == "FN_OCI_PIO" and not (r["fn_status"] or r["oci_status"] or foreign):
            f.append(finding("CON-008", "FN/OCI/PIO quota used for a resident Indian student"))
        lateral = r["admission_type"] == "LATERAL" or r["quota"] == "LATERAL"
        if lateral and current_cohort and r["year_of_study"] != 2:
            f.append(finding("CON-005", f"Lateral entry recorded in year {r['year_of_study']}"))
        expected_year = config.ADMISSION_YEAR - r["admission_year"] + 1 + (1 if lateral else 0)
        if not current_cohort and r["year_of_study"] != expected_year:
            f.append(finding("CON-010", f"Admitted in {r['admission_year']}, so year {expected_year} is expected, "
                             f"not year {r['year_of_study']}", expected=expected_year))
        if adm and current_cohort and not (win_start <= adm <= win_end):
            f.append(finding("CON-006", f"Admitted on {adm}, window is {win_start} to {win_end}",
                             admission_date=str(adm)))
        if adm and not current_cohort and adm.year != r["admission_year"]:
            f.append(finding("CON-006", f"Admission date {adm} is not in admission year {r['admission_year']}",
                             admission_date=str(adm)))
        if dob and adm and prog:
            lo, hi = config.AGE_LIMITS.get(prog.level, (14, 60))
            age = age_on(dob, adm)
            if not lo <= age <= hi:
                f.append(finding("CON-007", f"Age {age} at admission; expected {lo}-{hi} for {prog.level}",
                                 age=age))

        # --- intake ledger (current-year admissions that would otherwise be accepted) ------
        sn_type = "LATERAL" if lateral else SUPERNUMERARY_BY_QUOTA.get(r["quota"])
        seat_type = "SUPERNUMERARY" if sn_type else "WITHIN_INTAKE"
        if prog and current_cohort and not is_update and not any(x["severity"] == "ERROR" for x in f):
            if sn_type:
                n, cap = ledger.supernumerary[(prog.id, sn_type)] + 1, ledger.cap(prog.id, sn_type)
                frac = config.sn_cap_fraction(sn_type, prog.level)
                if n <= cap:
                    ledger.supernumerary[(prog.id, sn_type)] = n
                else:
                    ledger.overflow[(prog.id, sn_type)] += 1
                    n += ledger.overflow[(prog.id, sn_type)] - 1
                    f.append(finding("INT-002", f"{sn_type} seat #{n} exceeds the cap of {cap} "
                                     f"({frac:.0%} of {prog.approved_intake} for {prog.level})",
                                     quota=sn_type, seat_number=n, cap=cap, cap_fraction=frac,
                                     cap_verified=config.sn_cap_verified(sn_type)))
                    if not config.sn_cap_verified(sn_type):
                        f[-1]["severity"] = "WARNING"
                        f[-1]["message"] += "; this cap is not yet confirmed in the handbook in force, so the row " \
                                            "is accepted and sent for review"
            else:
                n = ledger.within[prog.id] + 1
                if n <= prog.approved_intake:
                    ledger.within[prog.id] = n
                    if prog.id not in ledger.near_noted and n >= math.ceil(prog.approved_intake * config.NEAR_CAP_WARNING):
                        ledger.near_noted.add(prog.id)
                        f.append(finding("INT-003", f"Seat {n} of {prog.approved_intake}: {prog.name} is now "
                                         f"{n / prog.approved_intake:.0%} full", seat_number=n,
                                         approved_intake=prog.approved_intake))
                else:
                    ledger.overflow[(prog.id, None)] += 1
                    n += ledger.overflow[(prog.id, None)] - 1
                    f.append(finding("INT-001", f"This would be admission #{n} in {prog.name}, "
                                     f"approved intake is {prog.approved_intake}",
                                     seat_number=n, approved_intake=prog.approved_intake))

        sev = {x["severity"] for x in f}
        status = "ERROR" if "ERROR" in sev else "WARNING" if "WARNING" in sev else "OK"
        results.append({
            "row_no": row_no, "status": status, "data": {**r, "seat_type": seat_type,
                                                         "supernumerary_type": sn_type},
            "programme_id": prog.id if prog else None, "is_update": is_update,
            "update_enrolment_id": update_of.id if update_of is not None else None, "changes": changes,
            "findings": f,
        })
    return results
